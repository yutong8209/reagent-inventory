from datetime import date, datetime, time, timedelta
from io import BytesIO
from typing import Optional
from urllib.parse import quote

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import func, or_
from sqlalchemy.orm import Session

from db import get_db
from models import BatchStock, Instrument, Reagent, Transaction, WorkGroup
from schemas import ReportReagentOption, TransactionListItem, UsageSummaryItem

router = APIRouter()


def _clean_text(value: str | None) -> str:
    return (value or "").strip()


def _safe_ratio(value: int | None) -> int:
    return value if value and value > 0 else 1


def _parse_int_filter(*values: object) -> set[int]:
    parsed: set[int] = set()
    for value in values:
        if value is None:
            continue
        for part in str(value).split(","):
            cleaned = part.strip()
            if cleaned.isdigit():
                parsed.add(int(cleaned))
    return parsed


def _parse_type_filter(*values: object) -> set[str]:
    allowed = {"IN", "OUT", "CALIBRATION"}
    parsed: set[str] = set()
    for value in values:
        if value is None:
            continue
        for part in str(value).split(","):
            normalized = part.strip().upper()
            if normalized in allowed:
                parsed.add(normalized)
    return parsed


def _safe_min_unit(transaction: Transaction, batch_stock: BatchStock | None, reagent: Reagent) -> str:
    return (
        _clean_text(transaction.min_unit)
        or _clean_text(batch_stock.min_unit if batch_stock else None)
        or _clean_text(reagent.unit)
        or "最小单位"
    )


def _transaction_min_quantity(transaction: Transaction) -> int:
    if transaction.type == "IN":
        return transaction.quantity * _safe_ratio(transaction.package_to_min_qty)
    return transaction.quantity


def _normalize_date_range(start_date: date | None, end_date: date | None) -> tuple[date, date]:
    today = date.today()
    normalized_end = end_date or today
    normalized_start = start_date or (normalized_end - timedelta(days=29))
    if normalized_start > normalized_end:
        normalized_start, normalized_end = normalized_end, normalized_start
    return normalized_start, normalized_end


def _build_datetime_range(start_date: date | None, end_date: date | None) -> tuple[datetime | None, datetime | None]:
    if start_date is None and end_date is None:
        return None, None

    normalized_start, normalized_end = _normalize_date_range(start_date, end_date)
    start_dt = datetime.combine(normalized_start, time.min)
    end_dt = datetime.combine(normalized_end + timedelta(days=1), time.min)
    return start_dt, end_dt


def _build_transaction_query(
    db: Session,
    *,
    work_group_id: int | None = None,
    work_group_ids: object | None = None,
    instrument_id: int | None = None,
    instrument_ids: object | None = None,
    reagent_id: int | None = None,
    keyword: str | None = None,
    transaction_type: str | None = None,
    transaction_types: object | None = None,
    start_date: date | None = None,
    end_date: date | None = None,
):
    query = (
        db.query(Transaction, Reagent, Instrument, WorkGroup, BatchStock)
        .join(Reagent, Transaction.reagent_id == Reagent.id)
        .join(Instrument, Reagent.instrument_id == Instrument.id)
        .join(WorkGroup, Instrument.work_group_id == WorkGroup.id)
        .outerjoin(BatchStock, Transaction.batch_stock_id == BatchStock.id)
    )

    work_group_filter = _parse_int_filter(work_group_ids, work_group_id)
    instrument_filter = _parse_int_filter(instrument_ids, instrument_id)

    if work_group_filter:
        query = query.filter(WorkGroup.id.in_(work_group_filter))
    if instrument_filter:
        query = query.filter(Instrument.id.in_(instrument_filter))
    if reagent_id is not None:
        query = query.filter(Reagent.id == reagent_id)

    cleaned_keyword = _clean_text(keyword)
    if cleaned_keyword:
        like_pattern = f"%{cleaned_keyword}%"
        query = query.filter(
            or_(
                Reagent.name.ilike(like_pattern),
                Reagent.manufacturer.ilike(like_pattern),
                Instrument.name.ilike(like_pattern),
                WorkGroup.name.ilike(like_pattern),
            )
        )

    type_filter = _parse_type_filter(transaction_types, transaction_type)
    if type_filter:
        query = query.filter(Transaction.type.in_(type_filter))

    start_dt, end_dt = _build_datetime_range(start_date, end_date)
    if start_dt is not None:
        query = query.filter(Transaction.created_at >= start_dt)
    if end_dt is not None:
        query = query.filter(Transaction.created_at < end_dt)

    return query


def _serialize_transaction_row(row) -> TransactionListItem:
    transaction, reagent, instrument, work_group, batch_stock = row
    return TransactionListItem(
        id=transaction.id,
        created_at=transaction.created_at,
        type=transaction.type,
        work_group_id=work_group.id if work_group else None,
        work_group_name=work_group.name if work_group else "",
        instrument_id=instrument.id if instrument else None,
        instrument_name=instrument.name if instrument else "",
        reagent_id=reagent.id,
        reagent_name=reagent.name,
        manufacturer=reagent.manufacturer,
        batch_no=batch_stock.batch_no if batch_stock and batch_stock.batch_no else "",
        quantity_min=_transaction_min_quantity(transaction),
        min_unit=_safe_min_unit(transaction, batch_stock, reagent),
        operator=transaction.operator,
        source=transaction.source,
    )


def _build_usage_query(
    db: Session,
    *,
    work_group_id: int | None = None,
    work_group_ids: object | None = None,
    instrument_id: int | None = None,
    instrument_ids: object | None = None,
    reagent_id: int | None = None,
    keyword: str | None = None,
    start_date: date | None = None,
    end_date: date | None = None,
):
    query = (
        db.query(
            Reagent.id.label("reagent_id"),
            Reagent.name.label("reagent_name"),
            Reagent.manufacturer.label("manufacturer"),
            Reagent.unit.label("min_unit"),
            Instrument.id.label("instrument_id"),
            Instrument.name.label("instrument_name"),
            WorkGroup.id.label("work_group_id"),
            WorkGroup.name.label("work_group_name"),
            func.coalesce(func.sum(Transaction.quantity), 0).label("total_usage_min"),
        )
        .join(Reagent, Transaction.reagent_id == Reagent.id)
        .join(Instrument, Reagent.instrument_id == Instrument.id)
        .join(WorkGroup, Instrument.work_group_id == WorkGroup.id)
        .filter(Transaction.type == "OUT")
    )

    work_group_filter = _parse_int_filter(work_group_ids, work_group_id)
    instrument_filter = _parse_int_filter(instrument_ids, instrument_id)

    if work_group_filter:
        query = query.filter(WorkGroup.id.in_(work_group_filter))
    if instrument_filter:
        query = query.filter(Instrument.id.in_(instrument_filter))
    if reagent_id is not None:
        query = query.filter(Reagent.id == reagent_id)

    cleaned_keyword = _clean_text(keyword)
    if cleaned_keyword:
        like_pattern = f"%{cleaned_keyword}%"
        query = query.filter(Reagent.name.ilike(like_pattern))

    start_dt, end_dt = _build_datetime_range(start_date, end_date)
    if start_dt is not None:
        query = query.filter(Transaction.created_at >= start_dt)
    if end_dt is not None:
        query = query.filter(Transaction.created_at < end_dt)

    return (
        query.group_by(
            Reagent.id,
            Reagent.name,
            Reagent.manufacturer,
            Reagent.unit,
            Instrument.id,
            Instrument.name,
            WorkGroup.id,
            WorkGroup.name,
        )
        .order_by(func.coalesce(func.sum(Transaction.quantity), 0).desc(), Reagent.name.asc(), Reagent.id.asc())
    )


def _build_reagent_report_query(
    db: Session,
    *,
    work_group_id: int | None = None,
    work_group_ids: object | None = None,
    instrument_id: int | None = None,
    instrument_ids: object | None = None,
    keyword: str | None = None,
):
    query = (
        db.query(Reagent, Instrument, WorkGroup)
        .join(Instrument, Reagent.instrument_id == Instrument.id)
        .join(WorkGroup, Instrument.work_group_id == WorkGroup.id)
        .order_by(WorkGroup.name.asc(), Instrument.name.asc(), Reagent.name.asc(), Reagent.id.asc())
    )

    work_group_filter = _parse_int_filter(work_group_ids, work_group_id)
    instrument_filter = _parse_int_filter(instrument_ids, instrument_id)

    if work_group_filter:
        query = query.filter(WorkGroup.id.in_(work_group_filter))
    if instrument_filter:
        query = query.filter(Instrument.id.in_(instrument_filter))

    cleaned_keyword = _clean_text(keyword)
    if cleaned_keyword:
        like_pattern = f"%{cleaned_keyword}%"
        query = query.filter(Reagent.name.ilike(like_pattern))

    return query


def _usage_filename(start_date: date | None, end_date: date | None) -> str:
    normalized_start, normalized_end = _normalize_date_range(start_date, end_date)
    return f"用量_{normalized_start.isoformat()}_{normalized_end.isoformat()}.xlsx"


def _transactions_filename(start_date: date | None, end_date: date | None) -> str:
    normalized_start, normalized_end = _normalize_date_range(start_date, end_date)
    return f"流水_{normalized_start.isoformat()}_{normalized_end.isoformat()}.xlsx"


def _style_print_sheet(worksheet, *, widths: list[int], landscape: bool = True):
    header_fill = PatternFill(fill_type="solid", fgColor="E8EEF7")
    header_font = Font(bold=True, color="17304F")
    thin = Side(style="thin", color="D9E3EF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border

    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[worksheet.cell(row=1, column=index).column_letter].width = width

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.page_margins.left = 0.25
    worksheet.page_margins.right = 0.25
    worksheet.page_margins.top = 0.5
    worksheet.page_margins.bottom = 0.5
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    if landscape:
        worksheet.page_setup.orientation = "landscape"


def _create_workbook(
    sheet_name: str,
    headers: list[str],
    rows: list[list[object]],
    *,
    widths: list[int] | None = None,
    landscape: bool = True,
) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(headers)
    for row in rows:
        sheet.append(row)

    if widths:
        _style_print_sheet(sheet, widths=widths, landscape=landscape)

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def _excel_response(buffer: BytesIO, filename: str) -> StreamingResponse:
    quoted_filename = quote(filename)
    headers = {
        "Content-Disposition": f"attachment; filename*=UTF-8''{quoted_filename}",
    }
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


def _type_label(value: str) -> str:
    return {
        "IN": "入库",
        "OUT": "出库",
        "CALIBRATION": "校准",
    }.get(value, value)


@router.get("/transactions", response_model=list[TransactionListItem])
def list_transactions(
    work_group_id: Optional[int] = Query(None),
    work_group_ids: Optional[str] = Query(None),
    instrument_id: Optional[int] = Query(None),
    instrument_ids: Optional[str] = Query(None),
    reagent_id: Optional[int] = Query(None),
    keyword: Optional[str] = Query(None),
    type: Optional[str] = Query(None),
    types: Optional[str] = Query(None),
    start_date: Optional[date] = Query(None),
    end_date: Optional[date] = Query(None),
    limit: int = Query(default=20, ge=1, le=100),
    db: Session = Depends(get_db),
):
    rows = (
        _build_transaction_query(
            db,
            work_group_id=work_group_id,
            work_group_ids=work_group_ids,
            instrument_id=instrument_id,
            instrument_ids=instrument_ids,
            reagent_id=reagent_id,
            keyword=keyword,
            transaction_type=type,
            transaction_types=types,
            start_date=start_date,
            end_date=end_date,
        )
        .order_by(Transaction.created_at.desc(), Transaction.id.desc())
        .limit(limit)
        .all()
    )
    return [_serialize_transaction_row(row) for row in rows]


@router.get("/usage", response_model=list[UsageSummaryItem])
def list_usage(
    work_group_id: Optional[int] = Query(None),
    work_group_ids: Optional[str] = Query(None),
    instrument_id: Optional[int] = Query(None),
    instrument_ids: Optional[str] = Query(None),
    reagent_id: Optional[int] = Query(None),
    keyword: Optional[str] = Query(None),
    start_date: Optional[date] = Query(None),
    end_date: Optional[date] = Query(None),
    db: Session = Depends(get_db),
):
    normalized_start, normalized_end = _normalize_date_range(start_date, end_date)
    rows = _build_usage_query(
        db,
        work_group_id=work_group_id,
        work_group_ids=work_group_ids,
        instrument_id=instrument_id,
        instrument_ids=instrument_ids,
        reagent_id=reagent_id,
        keyword=keyword,
        start_date=normalized_start,
        end_date=normalized_end,
    ).all()

    return [
        UsageSummaryItem(
            reagent_id=row.reagent_id,
            reagent_name=row.reagent_name,
            manufacturer=row.manufacturer,
            work_group_id=row.work_group_id,
            work_group_name=row.work_group_name or "",
            instrument_id=row.instrument_id,
            instrument_name=row.instrument_name or "",
            start_date=normalized_start,
            end_date=normalized_end,
            total_usage_min=int(row.total_usage_min or 0),
            min_unit=_clean_text(row.min_unit) or "最小单位",
        )
        for row in rows
    ]


@router.get("/reporting/reagents", response_model=list[ReportReagentOption])
def list_reporting_reagents(
    work_group_id: Optional[int] = Query(None),
    work_group_ids: Optional[str] = Query(None),
    instrument_id: Optional[int] = Query(None),
    instrument_ids: Optional[str] = Query(None),
    keyword: Optional[str] = Query(None),
    db: Session = Depends(get_db),
):
    rows = _build_reagent_report_query(
        db,
        work_group_id=work_group_id,
        work_group_ids=work_group_ids,
        instrument_id=instrument_id,
        instrument_ids=instrument_ids,
        keyword=keyword,
    ).all()

    return [
        ReportReagentOption(
            id=reagent.id,
            name=reagent.name,
            manufacturer=reagent.manufacturer,
            work_group_id=work_group.id if work_group else None,
            work_group_name=work_group.name if work_group else "",
            instrument_id=instrument.id if instrument else None,
            instrument_name=instrument.name if instrument else "",
            min_unit=_clean_text(reagent.unit) or "最小单位",
            is_archived=bool(reagent.is_archived),
        )
        for reagent, instrument, work_group in rows
    ]


@router.get("/transactions/export")
def export_transactions(
    work_group_id: Optional[int] = Query(None),
    work_group_ids: Optional[str] = Query(None),
    instrument_id: Optional[int] = Query(None),
    instrument_ids: Optional[str] = Query(None),
    reagent_id: Optional[int] = Query(None),
    keyword: Optional[str] = Query(None),
    type: Optional[str] = Query(None),
    types: Optional[str] = Query(None),
    start_date: Optional[date] = Query(None),
    end_date: Optional[date] = Query(None),
    db: Session = Depends(get_db),
):
    items = [
        _serialize_transaction_row(row)
        for row in _build_transaction_query(
            db,
            work_group_id=work_group_id,
            work_group_ids=work_group_ids,
            instrument_id=instrument_id,
            instrument_ids=instrument_ids,
            reagent_id=reagent_id,
            keyword=keyword,
            transaction_type=type,
            transaction_types=types,
            start_date=start_date,
            end_date=end_date,
        )
        .order_by(Transaction.created_at.desc(), Transaction.id.desc())
        .all()
    ]
    rows = [
        [
            item.created_at.strftime("%Y-%m-%d %H:%M:%S"),
            _type_label(item.type),
            f"{item.work_group_name or '-'}\n{item.instrument_name or '-'}",
            f"{item.reagent_name}\n批号: {item.batch_no or '-'}",
            f"{item.quantity_min} {item.min_unit or ''}".strip(),
            f"{item.operator or '-'}\n{'扫码' if item.source == 'scan' else '手工'}",
        ]
        for item in items
    ]
    buffer = _create_workbook(
        "流水",
        ["流水时间", "类型", "工作组 / 仪器", "试剂 / 批号", "数量", "操作人 / 来源"],
        rows,
        widths=[20, 10, 24, 32, 14, 18],
    )
    return _excel_response(buffer, _transactions_filename(start_date, end_date))


@router.get("/usage/export")
def export_usage(
    work_group_id: Optional[int] = Query(None),
    work_group_ids: Optional[str] = Query(None),
    instrument_id: Optional[int] = Query(None),
    instrument_ids: Optional[str] = Query(None),
    reagent_id: Optional[int] = Query(None),
    keyword: Optional[str] = Query(None),
    start_date: Optional[date] = Query(None),
    end_date: Optional[date] = Query(None),
    db: Session = Depends(get_db),
):
    items = list_usage(
        work_group_id=work_group_id,
        work_group_ids=work_group_ids,
        instrument_id=instrument_id,
        instrument_ids=instrument_ids,
        reagent_id=reagent_id,
        keyword=keyword,
        start_date=start_date,
        end_date=end_date,
        db=db,
    )
    rows = [
        [
            f"{item.work_group_name or '-'}\n{item.instrument_name or '-'}",
            f"{item.reagent_name}\n厂家: {item.manufacturer or '-'}",
            f"{item.start_date.isoformat()} 至\n{item.end_date.isoformat()}",
            item.total_usage_min,
            item.min_unit,
        ]
        for item in items
    ]
    buffer = _create_workbook(
        "用量",
        ["工作组 / 仪器", "试剂 / 厂家", "统计日期", "用量合计", "单位"],
        rows,
        widths=[26, 34, 22, 14, 12],
        landscape=False,
    )
    return _excel_response(buffer, _usage_filename(start_date, end_date))
