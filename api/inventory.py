from datetime import date, datetime, time, timedelta
from io import BytesIO
from typing import List, Optional
from urllib.parse import quote
from uuid import uuid4

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
import qrcode
from qrcode.image.svg import SvgPathImage
from sqlalchemy import func, or_, text
from sqlalchemy.orm import Session, aliased

from db import get_db
from models import BatchStock, CalibrationRecord, Instrument, Operator, ParseLog, Reagent, ReagentGroup, StockAdjustmentRecord, Transaction, UnitOption, WorkGroup
from schemas import (
    AdminPasswordRequest,
    BatchInfo,
    BatchStockAdjustmentRequest,
    StockAdjustmentRecordItem,
    StockAdjustmentRecordQuery,
    DashboardOverviewResponse,
    DashboardPendingTasks,
    DashboardWarningReagent,
    MobileOverviewResponse,
    InstrumentCatalogItem,
    InstrumentCreate,
    InstrumentInfo,
    InstrumentUpdate,
    InstrumentWarningDefaultsUpdate,
    ClearTransactionsRequest,
    OperatorCreate,
    OperatorResponse,
    OperatorUpdate,
    ReagentCatalogItem,
    ReagentGroupCreate,
    ReagentGroupResponse,
    ReagentGroupUpdate,
    ReagentInfo,
    ReagentUpdate,
    ReagentWarningSettingsUpdate,
    TransactionListItem,
    UnitOptionCreate,
    UnitOptionResponse,
    UnitOptionUpdate,
    WorkGroupCatalogItem,
    WorkGroupCreate,
    WorkGroupInfo,
    WorkGroupUpdate,
)

router = APIRouter()
ADMIN_ACTION_PASSWORD = "wtsyyjyk"

EXPORT_HEADERS = [
    "工作组",
    "仪器名称",
    "仪器型号",
    "试剂名称",
    "批号",
    "该批号库存最小单位数",
    "厂家",
    "试剂状态",
    "失效预警天数",
    "库存预警最小单位数",
    "总库存最小单位数",
    "最小单位",
    "最近效期",
    "距最近效期天数",
    "批次数",
    "最近出入库时间",
    "批号有效期",
    "该批号距失效天数",
    "是否过期",
    "包装单位",
    "每包装对应最小单位数",
    "整包装数量",
    "散装最小单位数量",
    "批次入库时间",
]
RED_FILL = PatternFill(fill_type="solid", fgColor="FDECEC")
YELLOW_FILL = PatternFill(fill_type="solid", fgColor="FFF6CF")
HEADER_FILL = PatternFill(fill_type="solid", fgColor="E8EEF7")
STATUS_RED_LEVELS = {"empty", "expired"}
STATUS_YELLOW_LEVELS = {"expiring", "low"}
STATUS_LABELS = {
    "empty": "零库存",
    "expired": "已过期",
    "expiring": "近效期",
    "low": "低库存",
    "normal": "正常",
}


def _clean_text(value: str | None) -> str:
    return (value or "").strip()


def _normalize_name(value: str | None) -> str:
    return "".join(_clean_text(value).lower().split())


def _is_calibrator_name(name: str | None) -> bool:
    return _clean_text(name).endswith("校准品")


def _safe_int(value: int | None, fallback: int) -> int:
    if value is None:
        return fallback
    return max(value, 0)


def _safe_ratio(value: int | None) -> int:
    return value if value and value > 0 else 1


def _unit_option_sort_key(option: UnitOption):
    scope_priority = {"both": 0, "package": 1, "min": 2}
    return (
        option.sort_order if option.sort_order is not None else 100,
        scope_priority.get(option.scope, 9),
        option.name.lower(),
        option.id,
    )


def _unit_option_matches_scope(option_scope: str, scope: str | None) -> bool:
    if not scope:
        return True
    if scope == "package":
        return option_scope in {"package", "both"}
    if scope == "min":
        return option_scope in {"min", "both"}
    return True


def _reagent_group_sort_key(group: ReagentGroup):
    return (
        group.sort_order if group.sort_order is not None else 100,
        group.name.lower(),
        group.id,
    )


def _serialize_reagent_group(group: ReagentGroup) -> ReagentGroupResponse:
    return ReagentGroupResponse.model_validate(group)


def _resolve_reagent_group(
    db: Session,
    reagent_group_id: int | None,
) -> ReagentGroup | None:
    if reagent_group_id is None:
        return None

    group = db.query(ReagentGroup).filter(ReagentGroup.id == reagent_group_id).first()
    if not group:
        raise HTTPException(status_code=400, detail="项目分组不存在")
    return group


def _require_admin_action_password(password: str):
    if password != ADMIN_ACTION_PASSWORD:
        raise HTTPException(status_code=403, detail="密码错误，不能执行危险操作")


def _days_until(target: date) -> int:
    return (target - date.today()).days


def _batch_total_min_units(batch: BatchStock) -> int:
    ratio = _safe_ratio(batch.package_to_min_qty)
    return (batch.quantity * ratio) + (batch.loose_min_quantity or 0)


def _set_batch_total_min_units(batch: BatchStock, total_min_units: int):
    ratio = _safe_ratio(batch.package_to_min_qty)
    normalized_total = max(int(total_min_units or 0), 0)
    batch.quantity = normalized_total // ratio
    batch.loose_min_quantity = normalized_total % ratio


def _reagent_total_min_units(db: Session, reagent_id: int) -> int:
    batches = (
        db.query(BatchStock)
        .filter(BatchStock.reagent_id == reagent_id, BatchStock.is_archived.is_(False))
        .all()
    )
    return sum(_batch_total_min_units(batch) for batch in batches)


def _effective_warning_values(reagent: Reagent) -> tuple[int, int, str]:
    instrument = reagent.instrument
    if reagent.follow_instrument and instrument:
        return (
            _safe_int(instrument.default_warning_days, 30),
            _safe_int(instrument.default_warning_stock, 5),
            "instrument",
        )

    return (
        _safe_int(reagent.warning_days, 30),
        _safe_int(reagent.warning_stock, 5),
        "reagent",
    )


def _effective_single_batch_warning_stock(reagent: Reagent, effective_warning_stock: int) -> int:
    if reagent.single_batch_warning_stock is not None:
        return _safe_int(reagent.single_batch_warning_stock, effective_warning_stock)
    instrument = reagent.instrument
    if instrument and instrument.default_single_batch_warning_stock is not None:
        return _safe_int(instrument.default_single_batch_warning_stock, effective_warning_stock)
    return effective_warning_stock


def _has_single_batch_warning(batch_items: list[BatchInfo], warning_stock: int) -> bool:
    threshold = max(_safe_int(warning_stock, 0), 0)
    return any(batch.available_min_quantity <= threshold for batch in batch_items)


def _build_status_level(
    *,
    has_expired_batch: bool,
    nearest_expiry_days: int | None,
    effective_warning_days: int,
    total_min_quantity: int,
    effective_warning_stock: int,
) -> str:
    if total_min_quantity <= 0:
        return "empty"
    if has_expired_batch:
        return "expired"
    if nearest_expiry_days is not None and nearest_expiry_days <= effective_warning_days:
        return "expiring"
    if total_min_quantity <= effective_warning_stock:
        return "low"
    return "normal"


def _apply_transaction_filters(query, filters: ClearTransactionsRequest):
    if filters.work_group_id is not None:
        query = query.join(Reagent, Transaction.reagent_id == Reagent.id)
        query = query.join(Instrument, Reagent.instrument_id == Instrument.id)
        query = query.filter(Instrument.work_group_id == filters.work_group_id)
        if filters.instrument_id is not None:
            query = query.filter(Instrument.id == filters.instrument_id)
    elif filters.instrument_id is not None:
        query = query.join(Reagent, Transaction.reagent_id == Reagent.id)
        query = query.filter(Reagent.instrument_id == filters.instrument_id)

    if filters.reagent_id is not None:
        query = query.filter(Transaction.reagent_id == filters.reagent_id)

    if filters.start_date is not None:
        query = query.filter(Transaction.created_at >= datetime.combine(filters.start_date, time.min))

    if filters.end_date is not None:
        query = query.filter(Transaction.created_at < datetime.combine(filters.end_date + timedelta(days=1), time.min))

    return query


def _apply_calibration_record_clear_filters(query, filters: ClearTransactionsRequest):
    source_reagent = aliased(Reagent)
    target_reagent = aliased(Reagent)
    source_instrument = aliased(Instrument)
    target_instrument = aliased(Instrument)

    query = (
        query
        .join(source_reagent, CalibrationRecord.calibrator_reagent_id == source_reagent.id)
        .join(target_reagent, CalibrationRecord.target_reagent_id == target_reagent.id)
    )

    needs_instrument_scope = filters.work_group_id is not None or filters.instrument_id is not None
    if needs_instrument_scope:
        query = (
            query
            .join(source_instrument, source_reagent.instrument_id == source_instrument.id)
            .join(target_instrument, target_reagent.instrument_id == target_instrument.id)
        )

    if filters.work_group_id is not None:
        query = query.filter(
            or_(
                source_instrument.work_group_id == filters.work_group_id,
                target_instrument.work_group_id == filters.work_group_id,
            )
        )

    if filters.instrument_id is not None:
        query = query.filter(
            or_(
                source_reagent.instrument_id == filters.instrument_id,
                target_reagent.instrument_id == filters.instrument_id,
            )
        )

    if filters.reagent_id is not None:
        query = query.filter(
            or_(
                CalibrationRecord.calibrator_reagent_id == filters.reagent_id,
                CalibrationRecord.target_reagent_id == filters.reagent_id,
            )
        )

    if filters.start_date is not None:
        query = query.filter(CalibrationRecord.created_at >= datetime.combine(filters.start_date, time.min))

    if filters.end_date is not None:
        query = query.filter(CalibrationRecord.created_at < datetime.combine(filters.end_date + timedelta(days=1), time.min))

    return query


def _serialize_batch(batch: BatchStock) -> BatchInfo:
    return BatchInfo(
        id=batch.id,
        batch_no=batch.batch_no,
        expiry_date=batch.expiry_date.isoformat(),
        expiry_days_left=_days_until(batch.expiry_date),
        quantity=batch.quantity,
        loose_min_quantity=batch.loose_min_quantity or 0,
        available_min_quantity=_batch_total_min_units(batch),
        package_unit=batch.package_unit,
        min_unit=batch.min_unit,
        package_to_min_qty=batch.package_to_min_qty,
        is_archived=bool(batch.is_archived),
        is_expired=batch.expiry_date < date.today(),
        created_at=batch.created_at,
    )


def _serialize_reagent(reagent: Reagent) -> ReagentCatalogItem:
    instrument = reagent.instrument
    work_group = instrument.work_group if instrument else None
    reagent_group = reagent.reagent_group
    effective_warning_days, effective_warning_stock, warning_source = _effective_warning_values(reagent)
    effective_single_batch_warning_stock = _effective_single_batch_warning_stock(
        reagent, effective_warning_stock
    )
    return ReagentCatalogItem(
        id=reagent.id,
        instrument_id=reagent.instrument_id,
        instrument_name=instrument.name if instrument else "",
        work_group_id=work_group.id if work_group else 0,
        work_group_name=work_group.name if work_group else "",
        reagent_group_id=reagent_group.id if reagent_group else None,
        reagent_group_name=reagent_group.name if reagent_group else None,
        reagent_group_sort_order=reagent_group.sort_order if reagent_group else 9999,
        name=reagent.name,
        manufacturer=reagent.manufacturer,
        is_calibrator=bool(reagent.is_calibrator),
        calibration_result_unit=_clean_text(reagent.calibration_result_unit) or "nmmol/L",
        unit=reagent.unit,
        package_unit=reagent.package_unit,
        package_to_min_qty=reagent.package_to_min_qty,
        warning_days=reagent.warning_days,
        warning_stock=reagent.warning_stock,
        single_batch_warning_stock=reagent.single_batch_warning_stock,
        follow_instrument=reagent.follow_instrument,
        is_archived=bool(reagent.is_archived),
        effective_warning_days=effective_warning_days,
        effective_warning_stock=effective_warning_stock,
        effective_single_batch_warning_stock=effective_single_batch_warning_stock,
        warning_source=warning_source,
    )


def _build_catalog(db: Session) -> list[WorkGroupCatalogItem]:
    work_groups = db.query(WorkGroup).order_by(WorkGroup.id.asc()).all()
    instruments = (
        db.query(Instrument)
        .order_by(Instrument.work_group_id.asc(), Instrument.id.asc())
        .all()
    )

    instruments_by_group: dict[int, list[InstrumentCatalogItem]] = {}
    for instrument in instruments:
        instruments_by_group.setdefault(instrument.work_group_id, []).append(
            InstrumentCatalogItem(
                id=instrument.id,
                work_group_id=instrument.work_group_id,
                name=instrument.name,
                model=instrument.model,
                default_warning_days=_safe_int(instrument.default_warning_days, 30),
                default_warning_stock=_safe_int(instrument.default_warning_stock, 5),
            )
        )

    return [
        WorkGroupCatalogItem(
            id=work_group.id,
            name=work_group.name,
            instruments=instruments_by_group.get(work_group.id, []),
        )
        for work_group in work_groups
    ]


def _build_reagent_info(
    reagent: Reagent,
    batches: list[BatchStock],
    last_transaction_at,
) -> ReagentInfo:
    batch_items = [_serialize_batch(batch) for batch in batches]
    total_package_quantity = sum(batch.quantity for batch in batches)
    total_min_quantity = sum(batch.available_min_quantity for batch in batch_items)
    nearest_expiry = batch_items[0] if batch_items else None
    effective_warning_days, effective_warning_stock, warning_source = _effective_warning_values(reagent)
    effective_single_batch_warning_stock = _effective_single_batch_warning_stock(
        reagent, effective_warning_stock
    )
    has_expired_batch = any(batch.is_expired for batch in batch_items)
    nearest_expiry_days = nearest_expiry.expiry_days_left if nearest_expiry else None
    has_expiring_soon = (
        not has_expired_batch
        and nearest_expiry_days is not None
        and nearest_expiry_days <= effective_warning_days
    )
    is_low_stock = total_min_quantity <= effective_warning_stock
    reagent_group = reagent.reagent_group
    status_level = _build_status_level(
        has_expired_batch=has_expired_batch,
        nearest_expiry_days=nearest_expiry_days,
        effective_warning_days=effective_warning_days,
        total_min_quantity=total_min_quantity,
        effective_warning_stock=effective_warning_stock,
    )

    return ReagentInfo(
        id=reagent.id,
        instrument_id=reagent.instrument_id,
        reagent_group_id=reagent_group.id if reagent_group else None,
        reagent_group_name=reagent_group.name if reagent_group else None,
        reagent_group_sort_order=reagent_group.sort_order if reagent_group else 9999,
        name=reagent.name,
        unit=reagent.unit,
        package_unit=reagent.package_unit,
        package_to_min_qty=reagent.package_to_min_qty,
        manufacturer=reagent.manufacturer,
        is_calibrator=bool(reagent.is_calibrator),
        calibration_result_unit=_clean_text(reagent.calibration_result_unit) or "nmmol/L",
        warning_days=reagent.warning_days,
        warning_stock=reagent.warning_stock,
        single_batch_warning_stock=reagent.single_batch_warning_stock,
        effective_warning_days=effective_warning_days,
        effective_warning_stock=effective_warning_stock,
        effective_single_batch_warning_stock=effective_single_batch_warning_stock,
        warning_source=warning_source,
        follow_instrument=reagent.follow_instrument,
        total_package_quantity=total_package_quantity,
        total_min_quantity=total_min_quantity,
        batch_count=len(batch_items),
        nearest_expiry_date=nearest_expiry.expiry_date if nearest_expiry else None,
        nearest_expiry_days=nearest_expiry_days,
        has_expired_batch=has_expired_batch,
        has_expiring_soon=has_expiring_soon,
        is_low_stock=is_low_stock,
        status_level=status_level,
        last_transaction_at=last_transaction_at,
        batches=batch_items,
    )


def _status_sort_key(reagent: ReagentInfo):
    level_priority = {
        "empty": 0,
        "expired": 1,
        "expiring": 2,
        "low": 3,
        "normal": 4,
    }
    return (
        level_priority.get(reagent.status_level, 99),
        reagent.nearest_expiry_days if reagent.nearest_expiry_days is not None else 10**9,
        reagent.name,
        reagent.id,
    )


INFECTIOUS_REAGENT_ORDER = {
    "hbsag": 0,
    "hbsab": 1,
    "hbeag": 2,
    "hbeab": 3,
    "hbcab": 4,
    "hcv": 5,
    "syphilis": 6,
    "hiv": 7,
}


def _normalized_reagent_order_name(name: str | None) -> str:
    return "".join((name or "").lower().replace("质控品", "").split())


def _custom_reagent_order(reagent: ReagentInfo) -> int:
    group_name = reagent.reagent_group_name or ""
    if group_name not in {"传染病标志物", "MAGLUMI X8质控品"}:
        return 999
    return INFECTIOUS_REAGENT_ORDER.get(_normalized_reagent_order_name(reagent.name), 999)


def _reagent_grouped_sort_key(reagent: ReagentInfo):
    return (
        reagent.reagent_group_sort_order if reagent.reagent_group_sort_order is not None else 9999,
        reagent.reagent_group_name or "未分组",
        _custom_reagent_order(reagent),
        reagent.name,
        reagent.id,
    )


def _query_inventory_data(
    db: Session,
    *,
    work_group_id: int | None = None,
    instrument_id: int | None = None,
    reagent_sort_mode: str = "status",
) -> list[WorkGroupInfo]:
    work_groups_query = db.query(WorkGroup).order_by(WorkGroup.id.asc())
    if work_group_id is not None:
        work_groups_query = work_groups_query.filter(WorkGroup.id == work_group_id)
    work_groups = work_groups_query.all()

    instruments_query = db.query(Instrument).order_by(Instrument.work_group_id.asc(), Instrument.id.asc())
    if work_group_id is not None:
        instruments_query = instruments_query.filter(Instrument.work_group_id == work_group_id)
    if instrument_id is not None:
        instruments_query = instruments_query.filter(Instrument.id == instrument_id)

    instruments = instruments_query.all()
    instruments_by_group: dict[int, list[Instrument]] = {}
    for instrument in instruments:
        instruments_by_group.setdefault(instrument.work_group_id, []).append(instrument)

    last_transaction_rows = (
        db.query(Transaction.reagent_id, func.max(Transaction.created_at))
        .group_by(Transaction.reagent_id)
        .all()
    )
    last_transactions = {reagent_id: last_created_at for reagent_id, last_created_at in last_transaction_rows}

    result: list[WorkGroupInfo] = []
    for work_group in work_groups:
        group_instruments = []
        for instrument in instruments_by_group.get(work_group.id, []):
            reagents = (
                db.query(Reagent)
                .filter(
                    Reagent.instrument_id == instrument.id,
                    Reagent.is_archived.is_(False),
                )
                .order_by(Reagent.name.asc(), Reagent.id.asc())
                .all()
            )

            reagent_items: list[ReagentInfo] = []
            for reagent in reagents:
                batches = (
                    db.query(BatchStock)
                    .filter(
                        BatchStock.reagent_id == reagent.id,
                        BatchStock.is_archived.is_(False),
                    )
                    .order_by(BatchStock.expiry_date.asc(), BatchStock.created_at.asc(), BatchStock.id.asc())
                    .all()
                )
                if not batches:
                    continue

                reagent_items.append(
                    _build_reagent_info(
                        reagent,
                        batches,
                        last_transactions.get(reagent.id),
                    )
                )

            sort_key = _reagent_grouped_sort_key if reagent_sort_mode == "group" else _status_sort_key
            reagent_items.sort(key=sort_key)
            group_instruments.append(
                InstrumentInfo(
                    id=instrument.id,
                    name=instrument.name,
                    model=instrument.model,
                    default_warning_days=_safe_int(instrument.default_warning_days, 30),
                    default_warning_stock=_safe_int(instrument.default_warning_stock, 5),
                    default_single_batch_warning_stock=instrument.default_single_batch_warning_stock,
                    reagents=reagent_items,
                )
            )

        result.append(
            WorkGroupInfo(
                id=work_group.id,
                name=work_group.name,
                instruments=group_instruments,
            )
        )

    return result


def _build_dashboard_pending_tasks(inventory_data: list[WorkGroupInfo]) -> DashboardPendingTasks:
    expired_batch_count = 0
    expiring_reagent_count = 0
    low_stock_reagent_count = 0
    low_stock_batch_count = 0
    low_or_empty_reagent_count = 0
    zero_stock_unarchived_count = 0

    for group in inventory_data:
        for instrument in group.instruments:
            for reagent in instrument.reagents:
                expired_batch_count += sum(1 for batch in reagent.batches if batch.is_expired)
                if reagent.status_level == "expiring":
                    expiring_reagent_count += 1
                if reagent.status_level == "low":
                    low_stock_reagent_count += 1
                if reagent.status_level in {"low", "empty"}:
                    low_or_empty_reagent_count += 1
                if reagent.total_min_quantity <= 0:
                    zero_stock_unarchived_count += 1
                else:
                    zero_stock_unarchived_count += sum(
                        1
                        for batch in reagent.batches
                        if batch.available_min_quantity <= 0
                    )
                if _has_single_batch_warning(reagent.batches, reagent.effective_single_batch_warning_stock):
                    low_stock_batch_count += 1

    return DashboardPendingTasks(
        expired_batch_count=expired_batch_count,
        expiring_reagent_count=expiring_reagent_count,
        low_stock_reagent_count=low_stock_reagent_count,
        low_stock_batch_count=low_stock_batch_count,
        low_or_empty_reagent_count=low_or_empty_reagent_count,
        zero_stock_unarchived_count=zero_stock_unarchived_count,
    )


def _build_dashboard_warning_reagents(inventory_data: list[WorkGroupInfo]) -> list[DashboardWarningReagent]:
    items: list[DashboardWarningReagent] = []
    for group in inventory_data:
        for instrument in group.instruments:
            for reagent in instrument.reagents:
                has_single_batch_warning = _has_single_batch_warning(
                    reagent.batches,
                    reagent.effective_single_batch_warning_stock,
                )
                if reagent.status_level == "normal" and not has_single_batch_warning:
                    continue
                items.append(
                    DashboardWarningReagent(
                        reagent_id=reagent.id,
                        reagent_name=reagent.name,
                        work_group_id=group.id,
                        work_group_name=group.name,
                        instrument_id=instrument.id,
                        instrument_name=instrument.name,
                        total_min_quantity=reagent.total_min_quantity,
                        min_unit=reagent.unit,
                        nearest_expiry_date=reagent.nearest_expiry_date,
                        nearest_expiry_days=reagent.nearest_expiry_days,
                        status_level=reagent.status_level,
                    )
                )

    level_priority = {"expired": 0, "empty": 1, "expiring": 2, "low": 3, "normal": 4}
    items.sort(
        key=lambda item: (
            level_priority.get(item.status_level, 99),
            item.nearest_expiry_days if item.nearest_expiry_days is not None else 10**9,
            item.total_min_quantity,
            item.reagent_name,
            item.reagent_id,
        )
    )
    return items[:6]


def _serialize_transaction_list_item(row) -> TransactionListItem:
    transaction, reagent, instrument, work_group, batch_stock = row
    min_unit = (
        _clean_text(transaction.min_unit)
        or _clean_text(batch_stock.min_unit if batch_stock else None)
        or _clean_text(reagent.unit)
        or "最小单位"
    )
    ratio = transaction.package_to_min_qty if transaction.package_to_min_qty and transaction.package_to_min_qty > 0 else 1
    quantity_min = transaction.quantity * ratio if transaction.type == "IN" else transaction.quantity
    return TransactionListItem(
        id=transaction.id,
        created_at=transaction.created_at,
        type=transaction.type,
        work_group_id=work_group.id if work_group else None,
        work_group_name=work_group.name if work_group else "",
        instrument_id=instrument.id if instrument else None,
        instrument_name=instrument.name if instrument else "",
        reagent_id=reagent.id,
        reagent_name=reagent.name,
        manufacturer=reagent.manufacturer,
        batch_no=batch_stock.batch_no if batch_stock and batch_stock.batch_no else "",
        quantity_min=quantity_min,
        min_unit=min_unit,
        operator=transaction.operator,
        source=transaction.source,
    )


def _query_recent_outbounds(
    db: Session,
    limit: int = 20,
    *,
    instrument_id: int | None = None,
) -> list[TransactionListItem]:
    query = (
        db.query(Transaction, Reagent, Instrument, WorkGroup, BatchStock)
        .join(Reagent, Transaction.reagent_id == Reagent.id)
        .join(Instrument, Reagent.instrument_id == Instrument.id)
        .join(WorkGroup, Instrument.work_group_id == WorkGroup.id)
        .outerjoin(BatchStock, Transaction.batch_stock_id == BatchStock.id)
        .filter(Transaction.type == "OUT")
    )
    if instrument_id is not None:
        query = query.filter(Instrument.id == instrument_id)

    rows = (
        query.order_by(Transaction.created_at.desc(), Transaction.id.desc())
        .limit(limit)
        .all()
    )
    return [_serialize_transaction_list_item(row) for row in rows]


def _sanitize_sheet_title(name: str, used_titles: set[str]) -> str:
    cleaned = "".join("_" if ch in '[]:*?/\\' else ch for ch in (name or "工作组")).strip() or "工作组"
    base = cleaned[:31]
    title = base
    index = 1
    while title in used_titles:
        suffix = f"_{index}"
        title = f"{base[: max(0, 31 - len(suffix))]}{suffix}" or f"Sheet{index}"
        index += 1
    used_titles.add(title)
    return title


def _format_export_datetime(value) -> str:
    if not value:
        return ""
    if isinstance(value, str):
        return value
    return value.strftime("%Y-%m-%d %H:%M:%S")


def _export_scope_label(
    db: Session,
    *,
    work_group_id: int | None,
    instrument_id: int | None,
) -> str:
    if instrument_id is not None:
        instrument = db.query(Instrument).filter(Instrument.id == instrument_id).first()
        if instrument:
            return instrument.name
    if work_group_id is not None:
        work_group = db.query(WorkGroup).filter(WorkGroup.id == work_group_id).first()
        if work_group:
            return work_group.name
    return "全部工作组"


def _build_export_rows(group: WorkGroupInfo) -> list[list[str | int]]:
    rows: list[list[str | int]] = []
    instruments = sorted(group.instruments, key=lambda item: (item.name, item.id))
    for instrument in instruments:
        reagents = sorted(
            instrument.reagents,
            key=lambda item: (
                instrument.name,
                _reagent_grouped_sort_key(item),
            ),
        )
        for reagent in reagents:
            for batch in reagent.batches:
                rows.append([
                    group.name,
                    instrument.name,
                    instrument.model or "",
                    reagent.name,
                    batch.batch_no,
                    batch.available_min_quantity,
                    reagent.manufacturer or "",
                    STATUS_LABELS.get(reagent.status_level, reagent.status_level),
                    reagent.effective_warning_days,
                    reagent.effective_warning_stock,
                    reagent.total_min_quantity,
                    reagent.unit,
                    reagent.nearest_expiry_date or "",
                    reagent.nearest_expiry_days if reagent.nearest_expiry_days is not None else "",
                    reagent.batch_count,
                    _format_export_datetime(reagent.last_transaction_at),
                    batch.expiry_date,
                    batch.expiry_days_left,
                    "是" if batch.is_expired else "否",
                    batch.package_unit,
                    batch.package_to_min_qty,
                    batch.quantity,
                    batch.loose_min_quantity,
                    _format_export_datetime(batch.created_at),
                ])
    return rows


def _export_reagent_category(reagent: ReagentInfo) -> str:
    name = _clean_text(reagent.name)
    if name.endswith("质控品"):
        return "qc"
    if reagent.is_calibrator or name.endswith("校准品"):
        return "calibrator"
    return "reagent"


def _apply_export_styles(worksheet):
    header_font = Font(bold=True)
    for cell in worksheet[1]:
        cell.font = header_font
        cell.fill = HEADER_FILL
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    status_col = EXPORT_HEADERS.index("试剂状态") + 1
    for row_index in range(2, worksheet.max_row + 1):
        status_value = worksheet.cell(row=row_index, column=status_col).value
        if status_value in {STATUS_LABELS[level] for level in STATUS_RED_LEVELS}:
            worksheet.cell(row=row_index, column=status_col).fill = RED_FILL
        elif status_value in {STATUS_LABELS[level] for level in STATUS_YELLOW_LEVELS}:
            worksheet.cell(row=row_index, column=status_col).fill = YELLOW_FILL

    for column_cells in worksheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 10), 28)


def _build_inventory_export_workbook(
    db: Session,
    *,
    work_group_id: int | None,
    instrument_id: int | None,
) -> tuple[BytesIO, str]:
    inventory_data = _query_inventory_data(
        db,
        work_group_id=work_group_id,
        instrument_id=instrument_id,
    )
    workbook = Workbook()
    workbook.remove(workbook.active)
    used_titles: set[str] = set()
    category_sheets = [
        ("reagent", "试剂库存"),
        ("calibrator", "校准品库存"),
        ("qc", "质控品库存"),
    ]
    rows_by_category: dict[str, list[list[str | int]]] = {
        category: [] for category, _ in category_sheets
    }

    for group in inventory_data:
        for instrument in sorted(group.instruments, key=lambda item: (item.name, item.id)):
            reagents = sorted(
                instrument.reagents,
                key=lambda item: (
                    instrument.name,
                    _reagent_grouped_sort_key(item),
                ),
            )
            for reagent in reagents:
                rows_by_category[_export_reagent_category(reagent)].extend(
                    _build_export_rows(
                        WorkGroupInfo(
                            id=group.id,
                            name=group.name,
                            instruments=[
                                InstrumentInfo(
                                    id=instrument.id,
                                    name=instrument.name,
                                    model=instrument.model,
                                    default_warning_days=instrument.default_warning_days,
                                    default_warning_stock=instrument.default_warning_stock,
                                    default_single_batch_warning_stock=instrument.default_single_batch_warning_stock,
                                    reagents=[reagent],
                                )
                            ],
                        )
                    )
                )

    for category, sheet_title in category_sheets:
        sheet = workbook.create_sheet(title=_sanitize_sheet_title(sheet_title, used_titles))
        sheet.append(EXPORT_HEADERS)
        for row in rows_by_category[category]:
            sheet.append(row)
        _apply_export_styles(sheet)

    if not workbook.sheetnames:
        sheet = workbook.create_sheet(title="库存")
        sheet.append(EXPORT_HEADERS)
        _apply_export_styles(sheet)

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    scope_label = _export_scope_label(db, work_group_id=work_group_id, instrument_id=instrument_id)
    timestamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    filename = f"库存导出_{scope_label}_{timestamp}.xlsx"
    return buffer, filename


def _validate_reagent_payload(
    *,
    db: Session,
    instrument_id: int,
    name: str,
    exclude_reagent_id: int | None = None,
) -> tuple[Instrument, str]:
    instrument = db.query(Instrument).filter(Instrument.id == instrument_id).first()
    if not instrument:
        raise HTTPException(status_code=400, detail=f"仪器 ID {instrument_id} 不存在")

    cleaned_name = _clean_text(name)
    if not cleaned_name:
        raise HTTPException(status_code=400, detail="试剂名称不能为空")

    normalized_name = _normalize_name(cleaned_name)
    existing = db.query(Reagent).filter(Reagent.instrument_id == instrument_id).all()
    for reagent in existing:
        if exclude_reagent_id is not None and reagent.id == exclude_reagent_id:
            continue
        if _normalize_name(reagent.name) == normalized_name:
            raise HTTPException(
                status_code=400,
                detail=f"仪器“{instrument.name}”下已存在同名试剂“{reagent.name}”",
            )

    return instrument, cleaned_name


@router.get("/inventory", response_model=List[WorkGroupInfo])
def get_inventory(
    instrument_id: Optional[int] = None,
    db: Session = Depends(get_db),
):
    return _query_inventory_data(db, instrument_id=instrument_id)


@router.get("/dashboard/overview", response_model=DashboardOverviewResponse)
def get_dashboard_overview(db: Session = Depends(get_db)):
    inventory_data = _query_inventory_data(db)
    return DashboardOverviewResponse(
        pending_tasks=_build_dashboard_pending_tasks(inventory_data),
        warning_reagents=_build_dashboard_warning_reagents(inventory_data),
        recent_outbounds=_query_recent_outbounds(db, limit=20),
    )


@router.get("/mobile/overview", response_model=MobileOverviewResponse)
def get_mobile_overview(
    instrument_id: int | None = Query(None, ge=1),
    db: Session = Depends(get_db),
):
    inventory_data = _query_inventory_data(db, instrument_id=instrument_id)
    return MobileOverviewResponse(
        pending_tasks=_build_dashboard_pending_tasks(inventory_data),
        warning_reagents=_build_dashboard_warning_reagents(inventory_data)[:4],
        recent_outbounds=_query_recent_outbounds(db, limit=5, instrument_id=instrument_id),
    )


@router.get("/mobile/qr")
def get_mobile_qr(
    target: str = Query(..., min_length=1, max_length=512),
):
    qr = qrcode.QRCode(
        version=None,
        error_correction=qrcode.constants.ERROR_CORRECT_M,
        box_size=8,
        border=1,
    )
    try:
        qr.add_data(target)
        qr.make(fit=True)
    except qrcode.exceptions.DataOverflowError:
        raise HTTPException(status_code=400, detail="二维码内容过长")

    image = qr.make_image(image_factory=SvgPathImage)
    buffer = BytesIO()
    image.save(buffer)
    svg = buffer.getvalue()
    return Response(content=svg, media_type="image/svg+xml")


@router.get("/inventory/export")
def export_inventory(
    work_group_id: int | None = Query(None),
    instrument_id: int | None = Query(None),
    db: Session = Depends(get_db),
):
    if instrument_id is not None:
        instrument = db.query(Instrument).filter(Instrument.id == instrument_id).first()
        if not instrument:
            raise HTTPException(status_code=404, detail="仪器不存在")
        if work_group_id is not None and instrument.work_group_id != work_group_id:
            raise HTTPException(status_code=400, detail="所选仪器不属于当前工作组")
        work_group_id = instrument.work_group_id
    elif work_group_id is not None:
        work_group = db.query(WorkGroup).filter(WorkGroup.id == work_group_id).first()
        if not work_group:
            raise HTTPException(status_code=404, detail="工作组不存在")

    buffer, filename = _build_inventory_export_workbook(
        db,
        work_group_id=work_group_id,
        instrument_id=instrument_id,
    )
    headers = {
        "Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}",
    }
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


@router.get("/work-groups")
def get_work_groups(db: Session = Depends(get_db)):
    work_groups = db.query(WorkGroup).order_by(WorkGroup.id.asc()).all()
    return [{"id": work_group.id, "name": work_group.name} for work_group in work_groups]


@router.get("/catalog/work-groups", response_model=List[WorkGroupCatalogItem])
def get_work_group_catalog(db: Session = Depends(get_db)):
    return _build_catalog(db)


@router.post("/work-groups", response_model=WorkGroupCatalogItem, status_code=201)
def create_work_group(data: WorkGroupCreate, db: Session = Depends(get_db)):
    name = _clean_text(data.name)
    if not name:
        raise HTTPException(status_code=400, detail="工作组名称不能为空")

    existing = db.query(WorkGroup).filter(WorkGroup.name == name).first()
    if existing:
        raise HTTPException(status_code=400, detail=f"工作组“{name}”已存在")

    work_group = WorkGroup(name=name)
    db.add(work_group)
    db.commit()
    db.refresh(work_group)

    return WorkGroupCatalogItem(id=work_group.id, name=work_group.name, instruments=[])


@router.put("/work-groups/{work_group_id}", response_model=WorkGroupCatalogItem)
def update_work_group(work_group_id: int, data: WorkGroupUpdate, db: Session = Depends(get_db)):
    work_group = db.query(WorkGroup).filter(WorkGroup.id == work_group_id).first()
    if not work_group:
        raise HTTPException(status_code=404, detail="工作组不存在")

    name = _clean_text(data.name)
    if not name:
        raise HTTPException(status_code=400, detail="工作组名称不能为空")

    existing = (
        db.query(WorkGroup)
        .filter(WorkGroup.name == name, WorkGroup.id != work_group_id)
        .first()
    )
    if existing:
        raise HTTPException(status_code=400, detail=f"工作组“{name}”已存在")

    work_group.name = name
    db.commit()
    db.refresh(work_group)

    instruments = (
        db.query(Instrument)
        .filter(Instrument.work_group_id == work_group.id)
        .order_by(Instrument.id.asc())
        .all()
    )
    return WorkGroupCatalogItem(
        id=work_group.id,
        name=work_group.name,
        instruments=[
            InstrumentCatalogItem(
                id=instrument.id,
                work_group_id=instrument.work_group_id,
                name=instrument.name,
                model=instrument.model,
                default_warning_days=_safe_int(instrument.default_warning_days, 30),
                default_warning_stock=_safe_int(instrument.default_warning_stock, 5),
            )
            for instrument in instruments
        ],
    )


@router.delete("/work-groups/{work_group_id}")
def delete_work_group(work_group_id: int, db: Session = Depends(get_db)):
    work_group = db.query(WorkGroup).filter(WorkGroup.id == work_group_id).first()
    if not work_group:
        raise HTTPException(status_code=404, detail="工作组不存在")

    instrument_count = db.query(Instrument).filter(Instrument.work_group_id == work_group_id).count()
    if instrument_count > 0:
        raise HTTPException(status_code=400, detail=f"当前工作组下还有 {instrument_count} 台仪器，请先删除或转移后再删除")

    db.delete(work_group)
    db.commit()
    return {"success": True, "work_group_id": work_group_id}


@router.get("/instruments")
def get_instruments(
    work_group_id: Optional[int] = None,
    db: Session = Depends(get_db),
):
    query = db.query(Instrument).order_by(Instrument.work_group_id.asc(), Instrument.id.asc())
    if work_group_id:
        query = query.filter(Instrument.work_group_id == work_group_id)

    instruments = query.all()
    return [
        {
            "id": instrument.id,
            "name": instrument.name,
            "model": instrument.model,
            "work_group_id": instrument.work_group_id,
            "default_warning_days": _safe_int(instrument.default_warning_days, 30),
            "default_warning_stock": _safe_int(instrument.default_warning_stock, 5),
            "default_single_batch_warning_stock": instrument.default_single_batch_warning_stock,
        }
        for instrument in instruments
    ]


@router.post("/instruments", response_model=InstrumentCatalogItem, status_code=201)
def create_instrument(data: InstrumentCreate, db: Session = Depends(get_db)):
    work_group = db.query(WorkGroup).filter(WorkGroup.id == data.work_group_id).first()
    if not work_group:
        raise HTTPException(status_code=400, detail="所属工作组不存在")

    name = _clean_text(data.name)
    if not name:
        raise HTTPException(status_code=400, detail="仪器名称不能为空")

    existing = (
        db.query(Instrument)
        .filter(Instrument.work_group_id == data.work_group_id, Instrument.name == name)
        .first()
    )
    if existing:
        raise HTTPException(status_code=400, detail=f"工作组“{work_group.name}”下已存在仪器“{name}”")

    instrument = Instrument(
        work_group_id=data.work_group_id,
        name=name,
        model=_clean_text(data.model) or None,
        default_warning_days=_safe_int(data.default_warning_days, 30),
        default_warning_stock=_safe_int(data.default_warning_stock, 5),
        default_single_batch_warning_stock=(
            _safe_int(data.default_single_batch_warning_stock, data.default_warning_stock)
            if data.default_single_batch_warning_stock is not None
            else None
        ),
    )
    db.add(instrument)
    db.commit()
    db.refresh(instrument)

    return InstrumentCatalogItem(
        id=instrument.id,
        work_group_id=instrument.work_group_id,
        name=instrument.name,
        model=instrument.model,
        default_warning_days=_safe_int(instrument.default_warning_days, 30),
        default_warning_stock=_safe_int(instrument.default_warning_stock, 5),
        default_single_batch_warning_stock=instrument.default_single_batch_warning_stock,
    )


@router.delete("/instruments/{instrument_id}")
def delete_instrument(instrument_id: int, db: Session = Depends(get_db)):
    instrument = db.query(Instrument).filter(Instrument.id == instrument_id).first()
    if not instrument:
        raise HTTPException(status_code=404, detail="仪器不存在")

    reagent_count = db.query(Reagent).filter(Reagent.instrument_id == instrument_id).count()
    if reagent_count > 0:
        raise HTTPException(status_code=400, detail=f"当前仪器下还有 {reagent_count} 个试剂档案，请先处理试剂后再删除")

    db.delete(instrument)
    db.commit()
    return {"success": True, "instrument_id": instrument_id}


@router.put("/instruments/{instrument_id}", response_model=InstrumentCatalogItem)
def update_instrument(instrument_id: int, data: InstrumentUpdate, db: Session = Depends(get_db)):
    instrument = db.query(Instrument).filter(Instrument.id == instrument_id).first()
    if not instrument:
        raise HTTPException(status_code=404, detail="仪器不存在")

    work_group = db.query(WorkGroup).filter(WorkGroup.id == data.work_group_id).first()
    if not work_group:
        raise HTTPException(status_code=400, detail="所属工作组不存在")

    name = _clean_text(data.name)
    if not name:
        raise HTTPException(status_code=400, detail="仪器名称不能为空")

    existing = (
        db.query(Instrument)
        .filter(
            Instrument.work_group_id == data.work_group_id,
            Instrument.name == name,
            Instrument.id != instrument_id,
        )
        .first()
    )
    if existing:
        raise HTTPException(status_code=400, detail=f"工作组“{work_group.name}”下已存在仪器“{name}”")

    instrument.work_group_id = data.work_group_id
    instrument.name = name
    instrument.model = _clean_text(data.model) or None
    instrument.default_warning_days = _safe_int(data.default_warning_days, 30)
    instrument.default_warning_stock = _safe_int(data.default_warning_stock, 5)
    instrument.default_single_batch_warning_stock = (
        _safe_int(data.default_single_batch_warning_stock, data.default_warning_stock)
        if data.default_single_batch_warning_stock is not None
        else None
    )
    db.commit()
    db.refresh(instrument)

    return InstrumentCatalogItem(
        id=instrument.id,
        work_group_id=instrument.work_group_id,
        name=instrument.name,
        model=instrument.model,
        default_warning_days=_safe_int(instrument.default_warning_days, 30),
        default_warning_stock=_safe_int(instrument.default_warning_stock, 5),
        default_single_batch_warning_stock=instrument.default_single_batch_warning_stock,
    )


@router.put("/instruments/{instrument_id}/warning-defaults", response_model=InstrumentCatalogItem)
def update_instrument_warning_defaults(
    instrument_id: int,
    data: InstrumentWarningDefaultsUpdate,
    db: Session = Depends(get_db),
):
    instrument = db.query(Instrument).filter(Instrument.id == instrument_id).first()
    if not instrument:
        raise HTTPException(status_code=404, detail="仪器不存在")

    instrument.default_warning_days = _safe_int(data.default_warning_days, 30)
    instrument.default_warning_stock = _safe_int(data.default_warning_stock, 5)
    instrument.default_single_batch_warning_stock = (
        _safe_int(data.default_single_batch_warning_stock, data.default_warning_stock)
        if data.default_single_batch_warning_stock is not None
        else None
    )
    db.commit()
    db.refresh(instrument)

    return InstrumentCatalogItem(
        id=instrument.id,
        work_group_id=instrument.work_group_id,
        name=instrument.name,
        model=instrument.model,
        default_warning_days=_safe_int(instrument.default_warning_days, 30),
        default_warning_stock=_safe_int(instrument.default_warning_stock, 5),
        default_single_batch_warning_stock=instrument.default_single_batch_warning_stock,
    )


@router.post("/instruments/{instrument_id}/warning-defaults/apply")
def apply_instrument_warning_defaults(
    instrument_id: int,
    db: Session = Depends(get_db),
):
    instrument = db.query(Instrument).filter(Instrument.id == instrument_id).first()
    if not instrument:
        raise HTTPException(status_code=404, detail="仪器不存在")

    updated_count = (
        db.query(Reagent)
        .filter(Reagent.instrument_id == instrument_id)
        .update({Reagent.follow_instrument: True}, synchronize_session=False)
    )
    db.commit()
    return {"success": True, "updated_count": updated_count}


@router.post("/instruments/{instrument_id}/warning-batch-update")
def batch_update_instrument_warnings(
    instrument_id: int,
    data: InstrumentWarningDefaultsUpdate,
    db: Session = Depends(get_db),
):
    instrument = db.query(Instrument).filter(Instrument.id == instrument_id).first()
    if not instrument:
        raise HTTPException(status_code=404, detail="仪器不存在")

    warning_days = _safe_int(data.default_warning_days, 30)
    warning_stock = _safe_int(data.default_warning_stock, 5)
    single_batch_warning_stock = (
        _safe_int(data.default_single_batch_warning_stock, warning_stock)
        if data.default_single_batch_warning_stock is not None
        else None
    )
    instrument.default_warning_days = warning_days
    instrument.default_warning_stock = warning_stock
    instrument.default_single_batch_warning_stock = single_batch_warning_stock

    updated_count = (
        db.query(Reagent)
        .filter(Reagent.instrument_id == instrument_id)
        .update(
            {
                Reagent.warning_days: warning_days,
                Reagent.warning_stock: warning_stock,
                Reagent.single_batch_warning_stock: None,
                Reagent.follow_instrument: False,
            },
            synchronize_session=False,
        )
    )
    db.commit()

    return {
        "success": True,
        "updated_count": updated_count,
        "default_warning_days": warning_days,
        "default_warning_stock": warning_stock,
        "default_single_batch_warning_stock": single_batch_warning_stock,
    }


@router.get("/reagents", response_model=List[ReagentCatalogItem])
def get_reagents(
    work_group_id: Optional[int] = None,
    instrument_id: Optional[int] = None,
    keyword: Optional[str] = None,
    status: str = Query(default="active"),
    db: Session = Depends(get_db),
):
    query = (
        db.query(Reagent)
        .join(Instrument, Reagent.instrument_id == Instrument.id)
        .join(WorkGroup, Instrument.work_group_id == WorkGroup.id)
        .order_by(WorkGroup.id.asc(), Instrument.id.asc(), Reagent.name.asc(), Reagent.id.asc())
    )

    if status == "active":
        query = query.filter(Reagent.is_archived.is_(False))
    elif status == "archived":
        query = query.filter(Reagent.is_archived.is_(True))
    elif status != "all":
        raise HTTPException(status_code=400, detail="状态筛选仅支持 active、archived 或 all")

    if work_group_id:
        query = query.filter(Instrument.work_group_id == work_group_id)

    if instrument_id:
        query = query.filter(Reagent.instrument_id == instrument_id)

    cleaned_keyword = _clean_text(keyword)
    if cleaned_keyword:
        query = query.filter(func.lower(Reagent.name).like(f"%{cleaned_keyword.lower()}%"))

    return [_serialize_reagent(reagent) for reagent in query.all()]


@router.put("/reagents/{reagent_id}", response_model=ReagentCatalogItem)
def update_reagent(reagent_id: int, data: ReagentUpdate, db: Session = Depends(get_db)):
    reagent = db.query(Reagent).filter(Reagent.id == reagent_id).first()
    if not reagent:
        raise HTTPException(status_code=404, detail="试剂不存在")

    instrument, name = _validate_reagent_payload(
        db=db,
        instrument_id=data.instrument_id,
        name=data.name,
        exclude_reagent_id=reagent_id,
    )
    reagent_group = _resolve_reagent_group(db, data.reagent_group_id)

    reagent.instrument_id = instrument.id
    reagent.reagent_group_id = reagent_group.id if reagent_group else None
    reagent.name = name
    reagent.manufacturer = _clean_text(data.manufacturer) or None
    reagent.is_calibrator = bool(data.is_calibrator) or _is_calibrator_name(name)
    reagent.calibration_result_unit = _clean_text(data.calibration_result_unit) or "nmmol/L"
    reagent.unit = _clean_text(data.unit) or reagent.unit
    reagent.package_unit = _clean_text(data.package_unit) or reagent.package_unit
    reagent.package_to_min_qty = data.package_to_min_qty
    reagent.warning_days = _safe_int(data.warning_days, 30)
    reagent.warning_stock = _safe_int(data.warning_stock, 5)
    if "single_batch_warning_stock" in data.model_fields_set:
        reagent.single_batch_warning_stock = data.single_batch_warning_stock
    reagent.follow_instrument = bool(data.follow_instrument)
    (
        db.query(BatchStock)
        .filter(BatchStock.reagent_id == reagent.id)
        .update(
            {
                BatchStock.package_unit: reagent.package_unit,
                BatchStock.min_unit: reagent.unit,
                BatchStock.package_to_min_qty: reagent.package_to_min_qty,
            },
            synchronize_session=False,
        )
    )
    db.commit()
    db.refresh(reagent)

    return _serialize_reagent(reagent)


@router.put("/reagents/{reagent_id}/warning-settings", response_model=ReagentCatalogItem)
def update_reagent_warning_settings(
    reagent_id: int,
    data: ReagentWarningSettingsUpdate,
    db: Session = Depends(get_db),
):
    reagent = db.query(Reagent).filter(Reagent.id == reagent_id).first()
    if not reagent:
        raise HTTPException(status_code=404, detail="试剂不存在")

    reagent.follow_instrument = bool(data.follow_instrument)
    reagent.warning_days = _safe_int(data.warning_days, 30)
    reagent.warning_stock = _safe_int(data.warning_stock, 5)
    if "single_batch_warning_stock" in data.model_fields_set:
        reagent.single_batch_warning_stock = data.single_batch_warning_stock
    db.commit()
    db.refresh(reagent)

    return _serialize_reagent(reagent)


@router.post("/reagents/{reagent_id}/archive")
def archive_reagent(reagent_id: int, db: Session = Depends(get_db)):
    reagent = (
        db.query(Reagent)
        .filter(Reagent.id == reagent_id, Reagent.is_archived.is_(False))
        .first()
    )
    if not reagent:
        raise HTTPException(status_code=404, detail="试剂不存在")

    total_min_quantity = _reagent_total_min_units(db, reagent_id)
    if total_min_quantity > 0:
        raise HTTPException(status_code=400, detail="仅支持归档零库存试剂")

    reagent.is_archived = True
    db.commit()
    return {"success": True, "reagent_id": reagent_id}


@router.post("/reagents/{reagent_id}/restore")
def restore_reagent(reagent_id: int, db: Session = Depends(get_db)):
    reagent = (
        db.query(Reagent)
        .filter(Reagent.id == reagent_id, Reagent.is_archived.is_(True))
        .first()
    )
    if not reagent:
        raise HTTPException(status_code=404, detail="试剂不存在或未归档")

    reagent.is_archived = False
    db.commit()
    return {"success": True, "reagent_id": reagent_id}


@router.post("/reagents/archive-zero-stock")
def archive_zero_stock_reagents(
    instrument_id: int | None = None,
    db: Session = Depends(get_db),
):
    query = db.query(Reagent).filter(Reagent.is_archived.is_(False))
    if instrument_id is not None:
        query = query.filter(Reagent.instrument_id == instrument_id)
    reagents = query.all()
    archived_ids: list[int] = []
    archived_batch_ids: list[int] = []

    for reagent in reagents:
        if _reagent_total_min_units(db, reagent.id) <= 0:
            reagent.is_archived = True
            archived_ids.append(reagent.id)
            continue

        zero_batches = (
            db.query(BatchStock)
            .filter(
                BatchStock.reagent_id == reagent.id,
                BatchStock.is_archived.is_(False),
            )
            .all()
        )
        for batch in zero_batches:
            if _batch_total_min_units(batch) <= 0:
                batch.is_archived = True
                archived_batch_ids.append(batch.id)

    db.commit()
    return {
        "success": True,
        "archived_count": len(archived_ids),
        "archived_batch_count": len(archived_batch_ids),
        "reagent_ids": archived_ids,
        "batch_stock_ids": archived_batch_ids,
    }


@router.get("/reagent-groups", response_model=List[ReagentGroupResponse])
def get_reagent_groups(
    include_inactive: bool = Query(False),
    db: Session = Depends(get_db),
):
    query = db.query(ReagentGroup)
    if not include_inactive:
        query = query.filter(ReagentGroup.is_active.is_(True))
    groups = query.all()
    groups.sort(key=_reagent_group_sort_key)
    return [_serialize_reagent_group(group) for group in groups]


@router.post("/reagent-groups", response_model=ReagentGroupResponse, status_code=201)
def create_reagent_group(data: ReagentGroupCreate, db: Session = Depends(get_db)):
    name = _clean_text(data.name)
    if not name:
        raise HTTPException(status_code=400, detail="项目分组名称不能为空")

    existing = db.query(ReagentGroup).filter(func.lower(ReagentGroup.name) == name.lower()).first()
    if existing:
        raise HTTPException(status_code=400, detail=f"项目分组“{name}”已存在")

    group = ReagentGroup(
        name=name,
        sort_order=_safe_int(data.sort_order, 100),
        is_active=bool(data.is_active),
    )
    db.add(group)
    db.commit()
    db.refresh(group)
    return _serialize_reagent_group(group)


@router.put("/reagent-groups/{group_id}", response_model=ReagentGroupResponse)
def update_reagent_group(group_id: int, data: ReagentGroupUpdate, db: Session = Depends(get_db)):
    group = db.query(ReagentGroup).filter(ReagentGroup.id == group_id).first()
    if not group:
        raise HTTPException(status_code=404, detail="项目分组不存在")

    name = _clean_text(data.name)
    if not name:
        raise HTTPException(status_code=400, detail="项目分组名称不能为空")

    existing = (
        db.query(ReagentGroup)
        .filter(func.lower(ReagentGroup.name) == name.lower(), ReagentGroup.id != group_id)
        .first()
    )
    if existing:
        raise HTTPException(status_code=400, detail=f"项目分组“{name}”已存在")

    group.name = name
    group.sort_order = _safe_int(data.sort_order, 100)
    group.is_active = bool(data.is_active)
    db.commit()
    db.refresh(group)
    return _serialize_reagent_group(group)


@router.delete("/reagent-groups/{group_id}")
def delete_reagent_group(group_id: int, db: Session = Depends(get_db)):
    group = db.query(ReagentGroup).filter(ReagentGroup.id == group_id).first()
    if not group:
        raise HTTPException(status_code=404, detail="项目分组不存在")

    reagent_count = db.query(Reagent).filter(Reagent.reagent_group_id == group_id).count()
    if reagent_count > 0:
        raise HTTPException(status_code=400, detail=f"当前分组已被 {reagent_count} 个试剂使用，请先更改试剂分组后再删除")

    db.delete(group)
    db.commit()
    return {"success": True, "group_id": group_id}


@router.get("/unit-options", response_model=List[UnitOptionResponse])
def get_unit_options(
    scope: Optional[str] = None,
    include_inactive: bool = False,
    db: Session = Depends(get_db),
):
    cleaned_scope = _clean_text(scope).lower() or None
    if cleaned_scope not in {None, "package", "min", "all"}:
        raise HTTPException(status_code=400, detail="单位类型仅支持 package、min 或 all")

    query = db.query(UnitOption)
    if not include_inactive:
        query = query.filter(UnitOption.is_active.is_(True))

    options = query.all()
    if cleaned_scope and cleaned_scope != "all":
        options = [item for item in options if _unit_option_matches_scope(item.scope, cleaned_scope)]

    options.sort(key=_unit_option_sort_key)
    return options


@router.post("/unit-options", response_model=UnitOptionResponse, status_code=201)
def create_unit_option(data: UnitOptionCreate, db: Session = Depends(get_db)):
    name = _clean_text(data.name)
    if not name:
        raise HTTPException(status_code=400, detail="单位名称不能为空")

    existing = db.query(UnitOption).filter(func.lower(UnitOption.name) == name.lower()).first()
    if existing:
        raise HTTPException(status_code=400, detail=f"单位“{name}”已存在")

    option = UnitOption(
        name=name,
        scope=data.scope,
        is_active=data.is_active,
        sort_order=_safe_int(data.sort_order, 100),
    )
    db.add(option)
    db.commit()
    db.refresh(option)
    return option


@router.put("/unit-options/{option_id}", response_model=UnitOptionResponse)
def update_unit_option(option_id: int, data: UnitOptionUpdate, db: Session = Depends(get_db)):
    option = db.query(UnitOption).filter(UnitOption.id == option_id).first()
    if not option:
        raise HTTPException(status_code=404, detail="单位选项不存在")

    name = _clean_text(data.name)
    if not name:
        raise HTTPException(status_code=400, detail="单位名称不能为空")

    existing = (
        db.query(UnitOption)
        .filter(func.lower(UnitOption.name) == name.lower(), UnitOption.id != option_id)
        .first()
    )
    if existing:
        raise HTTPException(status_code=400, detail=f"单位“{name}”已存在")

    option.name = name
    option.scope = data.scope
    option.is_active = bool(data.is_active)
    option.sort_order = _safe_int(data.sort_order, 100)
    db.commit()
    db.refresh(option)
    return option


@router.delete("/unit-options/{option_id}")
def delete_unit_option(option_id: int, db: Session = Depends(get_db)):
    option = db.query(UnitOption).filter(UnitOption.id == option_id).first()
    if not option:
        raise HTTPException(status_code=404, detail="单位选项不存在")

    db.delete(option)
    db.commit()
    return {"success": True, "option_id": option_id}


@router.get("/operators", response_model=List[OperatorResponse])
def get_operators(db: Session = Depends(get_db)):
    return db.query(Operator).order_by(Operator.is_active.desc(), Operator.name.asc()).all()


@router.post("/operators", response_model=OperatorResponse, status_code=201)
def create_operator(data: OperatorCreate, db: Session = Depends(get_db)):
    name = _clean_text(data.name)
    if not name:
        raise HTTPException(status_code=400, detail="操作人名称不能为空")

    existing = db.query(Operator).filter(func.lower(Operator.name) == name.lower()).first()
    if existing:
        raise HTTPException(status_code=400, detail=f"操作人“{name}”已存在")

    operator = Operator(name=name, is_active=bool(data.is_active))
    db.add(operator)
    db.commit()
    db.refresh(operator)
    return operator


@router.put("/operators/{operator_id}", response_model=OperatorResponse)
def update_operator(operator_id: int, data: OperatorUpdate, db: Session = Depends(get_db)):
    operator = db.query(Operator).filter(Operator.id == operator_id).first()
    if not operator:
        raise HTTPException(status_code=404, detail="操作人不存在")

    name = _clean_text(data.name)
    if not name:
        raise HTTPException(status_code=400, detail="操作人名称不能为空")

    existing = (
        db.query(Operator)
        .filter(func.lower(Operator.name) == name.lower(), Operator.id != operator_id)
        .first()
    )
    if existing:
        raise HTTPException(status_code=400, detail=f"操作人“{name}”已存在")

    operator.name = name
    operator.is_active = data.is_active
    db.commit()
    db.refresh(operator)
    return operator


@router.delete("/operators/{operator_id}")
def delete_operator(operator_id: int, db: Session = Depends(get_db)):
    operator = db.query(Operator).filter(Operator.id == operator_id).first()
    if not operator:
        raise HTTPException(status_code=404, detail="操作人不存在")

    db.delete(operator)
    db.commit()
    return {"success": True, "operator_id": operator_id}


@router.post("/admin/stock-adjustments/access", response_model=List[WorkGroupInfo])
def get_stock_adjustment_inventory(data: AdminPasswordRequest, db: Session = Depends(get_db)):
    _require_admin_action_password(data.password)
    return _query_inventory_data(db, reagent_sort_mode="group")


@router.post("/admin/stock-adjustments/apply")
def apply_stock_adjustments(data: BatchStockAdjustmentRequest, db: Session = Depends(get_db)):
    _require_admin_action_password(data.password)
    normalized_items = {
        int(item.batch_stock_id): max(int(item.target_min_quantity), 0)
        for item in data.items
    }
    if not normalized_items:
        return {"success": True, "updated_batch_count": 0, "updated_reagent_count": 0}

    batches = (
        db.query(BatchStock)
        .join(Reagent, BatchStock.reagent_id == Reagent.id)
        .filter(
            BatchStock.id.in_(list(normalized_items.keys())),
            BatchStock.is_archived.is_(False),
            Reagent.is_archived.is_(False),
        )
        .all()
    )
    batch_map = {batch.id: batch for batch in batches}
    missing_ids = [batch_id for batch_id in normalized_items if batch_id not in batch_map]
    if missing_ids:
        missing_text = ", ".join(str(item) for item in missing_ids)
        raise HTTPException(status_code=404, detail=f"以下批次不存在或已归档：{missing_text}")

    updated_batch_count = 0
    updated_reagent_ids: set[int] = set()
    adjustment_group_id = str(uuid4())
    for batch_id, target_min_quantity in normalized_items.items():
        batch = batch_map[batch_id]
        old_min_quantity = _batch_total_min_units(batch)
        if old_min_quantity == target_min_quantity:
            continue
        reagent = batch.reagent
        instrument = reagent.instrument if reagent else None
        work_group = instrument.work_group if instrument else None
        db.add(
            StockAdjustmentRecord(
                adjustment_group_id=adjustment_group_id,
                work_group_id=work_group.id if work_group else 0,
                work_group_name=work_group.name if work_group else "",
                instrument_id=instrument.id if instrument else 0,
                instrument_name=instrument.name if instrument else "",
                reagent_id=reagent.id if reagent else batch.reagent_id,
                reagent_name=reagent.name if reagent else "",
                batch_stock_id=batch.id,
                batch_no=batch.batch_no,
                expiry_date=batch.expiry_date,
                old_min_quantity=old_min_quantity,
                new_min_quantity=target_min_quantity,
                delta_min_quantity=target_min_quantity - old_min_quantity,
                min_unit=_clean_text(batch.min_unit) or _clean_text(reagent.unit if reagent else None) or "最小单位",
                package_unit=_clean_text(batch.package_unit) or _clean_text(reagent.package_unit if reagent else None) or "包装",
                package_to_min_qty=_safe_ratio(batch.package_to_min_qty),
                operator="库存调整",
                source="stock_adjustment",
            )
        )
        _set_batch_total_min_units(batch, target_min_quantity)
        updated_batch_count += 1
        updated_reagent_ids.add(batch.reagent_id)

    if not updated_batch_count:
        return {"success": True, "updated_batch_count": 0, "updated_reagent_count": 0}

    db.commit()
    return {
        "success": True,
        "updated_batch_count": updated_batch_count,
        "updated_reagent_count": len(updated_reagent_ids),
    }


@router.post("/admin/stock-adjustments/records", response_model=List[StockAdjustmentRecordItem])
def list_stock_adjustment_records(data: StockAdjustmentRecordQuery, db: Session = Depends(get_db)):
    _require_admin_action_password(data.password)
    query = db.query(StockAdjustmentRecord)
    if data.work_group_id is not None:
        query = query.filter(StockAdjustmentRecord.work_group_id == data.work_group_id)
    if data.instrument_id is not None:
        query = query.filter(StockAdjustmentRecord.instrument_id == data.instrument_id)
    if data.reagent_id is not None:
        query = query.filter(StockAdjustmentRecord.reagent_id == data.reagent_id)
    return (
        query
        .order_by(StockAdjustmentRecord.created_at.desc(), StockAdjustmentRecord.id.desc())
        .limit(data.limit)
        .all()
    )


@router.post("/admin/reset-reagent-data")
def reset_reagent_data(data: AdminPasswordRequest, db: Session = Depends(get_db)):
    _require_admin_action_password(data.password)
    db.query(StockAdjustmentRecord).delete(synchronize_session=False)
    db.query(CalibrationRecord).delete(synchronize_session=False)
    db.query(Transaction).delete(synchronize_session=False)
    db.query(BatchStock).delete(synchronize_session=False)
    db.query(ParseLog).delete(synchronize_session=False)
    db.query(Reagent).delete(synchronize_session=False)
    db.commit()

    try:
        db.execute(
            text(
                "DELETE FROM sqlite_sequence WHERE name IN ('calibration_records', 'transactions', 'batch_stocks', 'parse_logs', 'reagents')"
            )
        )
        db.commit()
    except Exception:
        db.rollback()

    return {"success": True}


@router.post("/admin/clear-stock")
def clear_stock(data: AdminPasswordRequest, db: Session = Depends(get_db)):
    _require_admin_action_password(data.password)
    transaction_batch_links_cleared = (
        db.query(Transaction)
        .filter(Transaction.batch_stock_id.is_not(None))
        .update({Transaction.batch_stock_id: None}, synchronize_session=False)
    )
    cleared_count = db.query(BatchStock).delete(synchronize_session=False)
    db.commit()

    try:
        db.execute(text("DELETE FROM sqlite_sequence WHERE name = 'batch_stocks'"))
        db.commit()
    except Exception:
        db.rollback()

    return {
        "success": True,
        "cleared_batch_count": cleared_count,
        "cleared_transaction_batch_links": transaction_batch_links_cleared,
    }


@router.post("/admin/clear-transactions")
def clear_transactions(data: ClearTransactionsRequest, db: Session = Depends(get_db)):
    _require_admin_action_password(data.password)
    if data.start_date and data.end_date and data.start_date > data.end_date:
        raise HTTPException(status_code=400, detail="开始日期不能晚于结束日期")

    query = db.query(Transaction)
    query = _apply_transaction_filters(query, data)

    transaction_ids = [item.id for item in query.with_entities(Transaction.id).all()]
    calibration_record_ids = {
        item.id
        for item in _apply_calibration_record_clear_filters(
            db.query(CalibrationRecord),
            data,
        )
        .with_entities(CalibrationRecord.id)
        .all()
    }
    if transaction_ids:
        linked_record_ids = (
            db.query(CalibrationRecord.id)
            .filter(CalibrationRecord.transaction_id.in_(transaction_ids))
            .all()
        )
        calibration_record_ids.update(item.id for item in linked_record_ids)

    deleted_calibration_record_count = 0
    if calibration_record_ids:
        deleted_calibration_record_count = (
            db.query(CalibrationRecord)
            .filter(CalibrationRecord.id.in_(calibration_record_ids))
            .delete(synchronize_session=False)
        )

    if not transaction_ids:
        db.commit()
        return {
            "success": True,
            "deleted_count": 0,
            "deleted_calibration_record_count": deleted_calibration_record_count,
        }

    deleted_count = (
        db.query(Transaction)
        .filter(Transaction.id.in_(transaction_ids))
        .delete(synchronize_session=False)
    )
    db.commit()

    return {
        "success": True,
        "deleted_count": deleted_count,
        "deleted_calibration_record_count": deleted_calibration_record_count,
    }
