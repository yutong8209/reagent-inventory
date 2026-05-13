import json
from datetime import date, datetime, time, timedelta
from io import BytesIO
from typing import Optional
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import or_
from sqlalchemy.orm import Session, aliased

from db import get_db
from models import BatchStock, CalibrationRecord, Instrument, Operator, Reagent, Transaction, WorkGroup
from schemas import (
    CalibrationBatchCandidate,
    CalibrationRecordCreate,
    CalibrationRecordItem,
    CalibrationRecordUpdate,
    CalibrationScanPreviewResponse,
    CalibrationTargetCandidate,
    CalibrationTargetPreviewResponse,
)
from services import BarcodeParser

router = APIRouter()


def _clean_text(value: str | None) -> str:
    return (value or "").strip()


def _normalize_name(value: str | None) -> str:
    return "".join(_clean_text(value).lower().split())


def _safe_ratio(value: int | None) -> int:
    return value if value and value > 0 else 1


def _parse_int_filter(*values: object) -> set[int]:
    parsed: set[int] = set()
    for value in values:
        if value is None:
            continue
        for part in str(value).split(","):
            cleaned = part.strip()
            if cleaned.isdigit():
                parsed.add(int(cleaned))
    return parsed


def _normalize_date_range(start_date: date | None, end_date: date | None) -> tuple[date, date]:
    today = date.today()
    normalized_end = end_date or today
    normalized_start = start_date or (normalized_end - timedelta(days=29))
    if normalized_start > normalized_end:
        normalized_start, normalized_end = normalized_end, normalized_start
    return normalized_start, normalized_end


def _build_datetime_range(start_date: date | None, end_date: date | None) -> tuple[datetime | None, datetime | None]:
    if start_date is None and end_date is None:
        return None, None

    normalized_start, normalized_end = _normalize_date_range(start_date, end_date)
    start_dt = datetime.combine(normalized_start, time.min)
    end_dt = datetime.combine(normalized_end + timedelta(days=1), time.min)
    return start_dt, end_dt


def _days_until(target: date) -> int:
    return (target - date.today()).days


def _batch_total_min_units(batch: BatchStock) -> int:
    return (batch.quantity * _safe_ratio(batch.package_to_min_qty)) + (batch.loose_min_quantity or 0)


def _set_batch_total_min_units(batch: BatchStock, total_min_units: int):
    ratio = _safe_ratio(batch.package_to_min_qty)
    batch.quantity = total_min_units // ratio
    batch.loose_min_quantity = total_min_units % ratio


def _require_operator(operator_name: str, db: Session) -> str:
    cleaned_operator = _clean_text(operator_name)
    if not cleaned_operator:
        raise HTTPException(status_code=400, detail="操作人不能为空")

    operator = (
        db.query(Operator)
        .filter(Operator.name == cleaned_operator, Operator.is_active.is_(True))
        .first()
    )
    if not operator:
        raise HTTPException(status_code=400, detail=f"操作人“{cleaned_operator}”不存在或已停用")

    return cleaned_operator


def _strip_calibrator_suffix(name: str | None) -> str:
    cleaned = _clean_text(name)
    if cleaned.endswith("校准品"):
        return cleaned[: -len("校准品")].strip()
    return cleaned


def _serialize_batch_candidate(batch: BatchStock) -> CalibrationBatchCandidate:
    return CalibrationBatchCandidate(
        batch_stock_id=batch.id,
        batch_no=batch.batch_no,
        expiry_date=batch.expiry_date.isoformat(),
        expiry_days_left=_days_until(batch.expiry_date),
        available_min_quantity=_batch_total_min_units(batch),
        min_unit=_clean_text(batch.min_unit) or "次",
        created_at=batch.created_at,
    )


def _load_target_batch_candidates(db: Session, reagent_id: int) -> list[CalibrationBatchCandidate]:
    batches = (
        db.query(BatchStock)
        .filter(
            BatchStock.reagent_id == reagent_id,
            BatchStock.is_archived.is_(False),
        )
        .order_by(BatchStock.expiry_date.desc(), BatchStock.created_at.desc(), BatchStock.id.desc())
        .all()
    )
    return [_serialize_batch_candidate(batch) for batch in batches]


def _load_available_batches(
    db: Session,
    reagent_id: int,
    *,
    newest_first: bool = False,
) -> list[BatchStock]:
    order_by = (
        (BatchStock.expiry_date.desc(), BatchStock.created_at.desc(), BatchStock.id.desc())
        if newest_first
        else (BatchStock.expiry_date.asc(), BatchStock.created_at.asc(), BatchStock.id.asc())
    )
    return (
        db.query(BatchStock)
        .filter(
            BatchStock.reagent_id == reagent_id,
            BatchStock.is_archived.is_(False),
            or_(BatchStock.quantity > 0, BatchStock.loose_min_quantity > 0),
        )
        .order_by(*order_by)
        .all()
    )


def _build_target_candidates(db: Session, calibrator: Reagent) -> list[Reagent]:
    target_base_name = _normalize_name(_strip_calibrator_suffix(calibrator.name))
    query = (
        db.query(Reagent)
        .filter(
            Reagent.instrument_id == calibrator.instrument_id,
            Reagent.is_archived.is_(False),
            Reagent.is_calibrator.is_(False),
        )
        .order_by(Reagent.name.asc(), Reagent.id.asc())
    )
    reagents = query.all()
    if target_base_name:
        exact = [reagent for reagent in reagents if _normalize_name(reagent.name) == target_base_name]
        if exact:
            ordered = exact + [reagent for reagent in reagents if reagent.id not in {item.id for item in exact}]
            return ordered
    return reagents


def _serialize_target_candidate(reagent: Reagent) -> CalibrationTargetCandidate:
    instrument = reagent.instrument
    work_group = instrument.work_group if instrument else None
    reagent_group = reagent.reagent_group
    return CalibrationTargetCandidate(
        reagent_id=reagent.id,
        reagent_name=reagent.name,
        instrument_id=instrument.id if instrument else 0,
        instrument_name=instrument.name if instrument else "",
        work_group_id=work_group.id if work_group else 0,
        work_group_name=work_group.name if work_group else "",
        reagent_group_id=reagent_group.id if reagent_group else None,
        reagent_group_name=reagent_group.name if reagent_group else None,
        reagent_group_sort_order=reagent_group.sort_order if reagent_group else 9999,
        calibration_result_unit=_clean_text(reagent.calibration_result_unit) or "nmmol/L",
    )


def _resolve_default_batch_values(candidates: list[CalibrationBatchCandidate]) -> tuple[str | None, str | None]:
    if not candidates:
        return None, None
    new_batch_no = candidates[0].batch_no
    old_batch_no = candidates[1].batch_no if len(candidates) > 1 else None
    return old_batch_no, new_batch_no


def _resolve_scanned_reagent_batch_values(
    candidates: list[CalibrationBatchCandidate],
    scanned_batch_no: str | None,
) -> tuple[str | None, str | None]:
    cleaned_scanned_batch = _clean_text(scanned_batch_no)
    if not candidates or not cleaned_scanned_batch:
        return _resolve_default_batch_values(candidates)

    for index, candidate in enumerate(candidates):
        if candidate.batch_no != cleaned_scanned_batch:
            continue
        old_batch_no = candidates[index + 1].batch_no if index + 1 < len(candidates) else None
        return old_batch_no, candidate.batch_no

    return _resolve_default_batch_values(candidates)


def _build_target_preview_payload(target: Reagent, db: Session) -> CalibrationTargetPreviewResponse:
    batch_candidates = _load_target_batch_candidates(db, target.id)
    default_old_batch_no, default_new_batch_no = _resolve_default_batch_values(batch_candidates)
    return CalibrationTargetPreviewResponse(
        target_reagent_id=target.id,
        target_reagent_name=target.name,
        result_unit=_clean_text(target.calibration_result_unit) or "nmmol/L",
        old_batch_candidates=batch_candidates,
        new_batch_candidates=batch_candidates,
        default_old_batch_no=default_old_batch_no,
        default_new_batch_no=default_new_batch_no,
    )


def _match_calibrator_batch(
    db: Session,
    *,
    calibrator_id: int,
    batch_no: str | None,
    expiry_date: date | None,
) -> BatchStock | None:
    query = db.query(BatchStock).filter(
        BatchStock.reagent_id == calibrator_id,
        BatchStock.is_archived.is_(False),
    )
    cleaned_batch_no = _clean_text(batch_no)
    if cleaned_batch_no:
        query = query.filter(BatchStock.batch_no == cleaned_batch_no)
    if expiry_date is not None:
        query = query.filter(BatchStock.expiry_date == expiry_date)

    batch = query.order_by(BatchStock.created_at.desc(), BatchStock.id.desc()).first()
    if batch:
        return batch

    if cleaned_batch_no:
        return None

    return (
        db.query(BatchStock)
        .filter(
            BatchStock.reagent_id == calibrator_id,
            BatchStock.is_archived.is_(False),
        )
        .order_by(BatchStock.expiry_date.desc(), BatchStock.created_at.desc(), BatchStock.id.desc())
        .first()
    )


def _serialize_record(record: CalibrationRecord, db: Session) -> CalibrationRecordItem:
    calibrator = record.calibrator_reagent
    target = record.target_reagent
    if calibrator.is_calibrator:
        target_candidates = [_serialize_target_candidate(item) for item in _build_target_candidates(db, calibrator)]
    else:
        target_candidates = [_serialize_target_candidate(target)]
    batch_candidates = _load_target_batch_candidates(db, target.id)
    calibration_source = "calibrator" if calibrator.is_calibrator else "reagent"
    return CalibrationRecordItem(
        id=record.id,
        calibration_source=calibration_source,
        calibrator_reagent_id=record.calibrator_reagent_id,
        calibrator_name=calibrator.name,
        calibrator_batch_stock_id=record.calibrator_batch_stock_id,
        calibrator_batch_no=record.calibrator_batch_stock.batch_no,
        target_reagent_id=record.target_reagent_id,
        target_reagent_name=target.name,
        target_old_batch_no=record.target_old_batch_no,
        target_new_batch_no=record.target_new_batch_no,
        usage_times=record.usage_times,
        result_unit=record.result_unit,
        operator=record.operator,
        old_result_1=record.old_result_1,
        old_result_2=record.old_result_2,
        old_result_3=record.old_result_3,
        old_result_4=record.old_result_4,
        old_result_5=record.old_result_5,
        new_result_1=record.new_result_1,
        new_result_2=record.new_result_2,
        new_result_3=record.new_result_3,
        new_result_4=record.new_result_4,
        new_result_5=record.new_result_5,
        target_options=target_candidates,
        old_batch_candidates=batch_candidates,
        new_batch_candidates=batch_candidates,
        created_at=record.created_at,
        updated_at=record.updated_at,
    )


def _build_calibration_record_query(
    db: Session,
    *,
    work_group_ids: object | None = None,
    instrument_ids: object | None = None,
    keyword: str | None = None,
    source: str | None = None,
    start_date: date | None = None,
    end_date: date | None = None,
):
    source_reagent = aliased(Reagent)
    target_reagent = aliased(Reagent)
    source_batch = aliased(BatchStock)
    target_instrument = aliased(Instrument)
    target_work_group = aliased(WorkGroup)

    query = (
        db.query(CalibrationRecord)
        .join(source_reagent, CalibrationRecord.calibrator_reagent_id == source_reagent.id)
        .join(source_batch, CalibrationRecord.calibrator_batch_stock_id == source_batch.id)
        .join(target_reagent, CalibrationRecord.target_reagent_id == target_reagent.id)
        .join(target_instrument, target_reagent.instrument_id == target_instrument.id)
        .join(target_work_group, target_instrument.work_group_id == target_work_group.id)
    )

    work_group_filter = _parse_int_filter(work_group_ids)
    instrument_filter = _parse_int_filter(instrument_ids)
    if work_group_filter:
        query = query.filter(target_work_group.id.in_(work_group_filter))
    if instrument_filter:
        query = query.filter(target_instrument.id.in_(instrument_filter))

    normalized_source = _clean_text(source).lower()
    if normalized_source == "calibrator":
        query = query.filter(source_reagent.is_calibrator.is_(True))
    elif normalized_source == "reagent":
        query = query.filter(source_reagent.is_calibrator.is_(False))

    cleaned_keyword = _clean_text(keyword)
    if cleaned_keyword:
        like_pattern = f"%{cleaned_keyword}%"
        query = query.filter(
            or_(
                source_reagent.name.ilike(like_pattern),
                target_reagent.name.ilike(like_pattern),
                source_batch.batch_no.ilike(like_pattern),
                CalibrationRecord.target_old_batch_no.ilike(like_pattern),
                CalibrationRecord.target_new_batch_no.ilike(like_pattern),
                CalibrationRecord.operator.ilike(like_pattern),
            )
        )

    start_dt, end_dt = _build_datetime_range(start_date, end_date)
    if start_dt is not None:
        query = query.filter(CalibrationRecord.created_at >= start_dt)
    if end_dt is not None:
        query = query.filter(CalibrationRecord.created_at < end_dt)

    return query


def _calibration_source_label(record: CalibrationRecord) -> str:
    calibrator = record.calibrator_reagent
    return "独立校准品" if calibrator and calibrator.is_calibrator else "试剂自带校准"


def _calibration_filename(start_date: date | None, end_date: date | None) -> str:
    normalized_start, normalized_end = _normalize_date_range(start_date, end_date)
    return f"校准记录_{normalized_start.isoformat()}_{normalized_end.isoformat()}.xlsx"


def _excel_response(buffer: BytesIO, filename: str) -> StreamingResponse:
    quoted_filename = quote(filename)
    headers = {
        "Content-Disposition": f"attachment; filename*=UTF-8''{quoted_filename}",
    }
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


def _style_calibration_export_sheet(worksheet):
    header_fill = PatternFill(fill_type="solid", fgColor="DDEFEA")
    group_fill = PatternFill(fill_type="solid", fgColor="F8FCFF")
    header_font = Font(bold=True, color="17304F")
    thin = Side(style="thin", color="BFD8CF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    widths = [20, 18, 24, 14, 18, 22, 24, 12, 12, 12, 12, 12, 12, 14]

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row_index, row in enumerate(worksheet.iter_rows(min_row=2), start=2):
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
            if row_index % 2 == 0:
                cell.fill = group_fill

    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[worksheet.cell(row=1, column=index).column_letter].width = width

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.page_margins.left = 0.25
    worksheet.page_margins.right = 0.25
    worksheet.page_margins.top = 0.5
    worksheet.page_margins.bottom = 0.5
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True


def _format_calibration_result(value: str | None, unit: str) -> str:
    cleaned = _clean_text(value)
    if not cleaned:
        return ""
    if any(char.isalpha() for char in cleaned) or "/" in cleaned:
        return cleaned
    return f"{cleaned}{unit}"


def _build_calibration_export_workbook(records: list[CalibrationRecord], start_date: date | None, end_date: date | None) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "校准记录"
    sheet.append(
        [
            "记录时间",
            "被校准试剂",
            "工作组 / 仪器",
            "校准来源",
            "校准品批号",
            "校准来源试剂",
            "新旧批号",
            "结果1",
            "结果2",
            "结果3",
            "结果4",
            "结果5",
            "校准用量",
            "操作人",
        ]
    )

    for record in records:
        start_row = sheet.max_row + 1
        target = record.target_reagent
        calibrator = record.calibrator_reagent
        batch = record.calibrator_batch_stock
        instrument = target.instrument if target else None
        work_group = instrument.work_group if instrument else None
        unit = _clean_text(record.result_unit) or "nmmol/L"
        shared = [
            record.created_at.strftime("%Y-%m-%d %H:%M:%S"),
            target.name if target else "-",
            f"{work_group.name if work_group else '-'}\n{instrument.name if instrument else '-'}",
            _calibration_source_label(record),
            batch.batch_no if batch else "-",
            calibrator.name if calibrator else "-",
        ]
        usage_text = f"{record.usage_times}{batch.min_unit if batch and batch.min_unit else '次'}"
        sheet.append(
            shared
            + [
                f"旧批号: {record.target_old_batch_no or '-'}",
                _format_calibration_result(record.old_result_1, unit),
                _format_calibration_result(record.old_result_2, unit),
                _format_calibration_result(record.old_result_3, unit),
                _format_calibration_result(record.old_result_4, unit),
                _format_calibration_result(record.old_result_5, unit),
                usage_text,
                record.operator,
            ]
        )
        sheet.append(
            shared
            + [
                f"新批号: {record.target_new_batch_no or '-'}",
                _format_calibration_result(record.new_result_1, unit),
                _format_calibration_result(record.new_result_2, unit),
                _format_calibration_result(record.new_result_3, unit),
                _format_calibration_result(record.new_result_4, unit),
                _format_calibration_result(record.new_result_5, unit),
                usage_text,
                record.operator,
            ]
        )
        for column in range(1, 7):
            sheet.merge_cells(start_row=start_row, start_column=column, end_row=start_row + 1, end_column=column)
        for column in range(13, 15):
            sheet.merge_cells(start_row=start_row, start_column=column, end_row=start_row + 1, end_column=column)

    _style_calibration_export_sheet(sheet)
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def _calibration_transaction_remark(
    *,
    target_reagent_id: int,
    target_reagent_name: str,
    target_old_batch_no: str | None,
    target_new_batch_no: str | None,
    usage_times: int,
    result_unit: str,
) -> str:
    return json.dumps(
        {
            "target_reagent_id": target_reagent_id,
            "target_reagent_name": target_reagent_name,
            "target_old_batch_no": target_old_batch_no,
            "target_new_batch_no": target_new_batch_no,
            "usage_times": usage_times,
            "result_unit": result_unit,
        },
        ensure_ascii=False,
    )


def _create_linked_outbound_transaction(
    *,
    source_reagent: Reagent,
    source_batch: BatchStock,
    operator: str,
    raw_barcode: str | None,
    calibration_record_id: int,
    db: Session,
    outbound_reagent_id: int | None = None,
    outbound_batch_stock_id: int | None = None,
    outbound_quantity: int = 1,
    outbound_unit: str = "package",
):
    outbound_reagent = source_reagent
    if outbound_reagent_id:
        found = db.query(Reagent).filter(Reagent.id == outbound_reagent_id, Reagent.is_archived.is_(False)).first()
        if not found:
            raise HTTPException(status_code=400, detail="出库试剂不存在或已归档")
        outbound_reagent = found

    if outbound_batch_stock_id:
        target_batch = (
            db.query(BatchStock)
            .filter(
                BatchStock.id == outbound_batch_stock_id,
                BatchStock.reagent_id == outbound_reagent.id,
                BatchStock.is_archived.is_(False),
            )
            .first()
        )
        if not target_batch:
            raise HTTPException(status_code=400, detail="出库批号不存在或与试剂不匹配")
    else:
        target_batch = source_batch

    package_to_min_qty = _safe_ratio(target_batch.package_to_min_qty)
    if outbound_unit == "package":
        requested_min_quantity = outbound_quantity * package_to_min_qty
    else:
        requested_min_quantity = outbound_quantity

    available_min_quantity = _batch_total_min_units(target_batch)
    if available_min_quantity < requested_min_quantity:
        unit_label = target_batch.package_unit or "盒" if outbound_unit == "package" else target_batch.min_unit or "最小单位"
        raise HTTPException(
            status_code=400,
            detail=f"{outbound_reagent.name} 批号 {target_batch.batch_no} 库存不足，不能同时出库 {outbound_quantity}{unit_label}",
        )

    _set_batch_total_min_units(target_batch, available_min_quantity - requested_min_quantity)
    db.add(
        Transaction(
            reagent_id=outbound_reagent.id,
            batch_stock_id=target_batch.id,
            type="OUT",
            quantity=requested_min_quantity,
            operator=operator,
            source="scan",
            raw_barcode=_clean_text(raw_barcode) or f"CAL-OUT::{outbound_reagent.id}",
            package_unit=target_batch.package_unit,
            min_unit=target_batch.min_unit,
            package_to_min_qty=package_to_min_qty,
            remark=json.dumps(
                {
                    "from_calibration": True,
                    "calibration_record_id": calibration_record_id,
                    "outbound_unit": outbound_unit,
                    "requested_quantity": outbound_quantity,
                    "requested_min_quantity": requested_min_quantity,
                    "deduction_mode": "scanned_batch",
                    "scanned_batch_no": target_batch.batch_no,
                    "scanned_expiry_date": target_batch.expiry_date.isoformat()
                    if target_batch.expiry_date
                    else None,
                    "deducted_min_quantity": requested_min_quantity,
                },
                ensure_ascii=False,
            ),
        )
    )


@router.post("/calibration/scan-preview", response_model=CalibrationScanPreviewResponse)
def calibration_scan_preview(payload: dict, db: Session = Depends(get_db)):
    barcode = _clean_text(payload.get("barcode"))
    if not barcode:
        raise HTTPException(status_code=400, detail="请先扫描校准条码")

    parser = BarcodeParser(db)
    parsed = parser.parse(barcode)
    if not parsed.reagent_id:
        raise HTTPException(status_code=404, detail="当前条码未匹配到试剂档案，不能进入校准流程")

    scanned_reagent = db.query(Reagent).filter(Reagent.id == parsed.reagent_id).first()
    if not scanned_reagent or scanned_reagent.is_archived:
        raise HTTPException(status_code=404, detail="当前扫码试剂不存在或已归档")

    scanned_batch = _match_calibrator_batch(
        db,
        calibrator_id=scanned_reagent.id,
        batch_no=parsed.extracted_batch_no,
        expiry_date=parsed.extracted_expiry_date,
    )
    if not scanned_batch:
        raise HTTPException(status_code=404, detail="未找到当前扫码批号对应的库存批次")

    if scanned_reagent.is_calibrator:
        calibration_source = "calibrator"
        target_candidates = _build_target_candidates(db, scanned_reagent)
        target = target_candidates[0] if target_candidates else None
        batch_candidates = _load_target_batch_candidates(db, target.id) if target else []
        default_old_batch_no, default_new_batch_no = _resolve_default_batch_values(batch_candidates)
    else:
        calibration_source = "reagent"
        target_candidates = [scanned_reagent]
        target = scanned_reagent
        batch_candidates = _load_target_batch_candidates(db, scanned_reagent.id)
        default_old_batch_no, default_new_batch_no = _resolve_scanned_reagent_batch_values(
            batch_candidates,
            scanned_batch.batch_no,
        )
    target_options = [_serialize_target_candidate(item) for item in target_candidates]

    return CalibrationScanPreviewResponse(
        calibration_source=calibration_source,
        calibrator_reagent_id=scanned_reagent.id,
        calibrator_name=scanned_reagent.name,
        calibrator_batch_stock_id=scanned_batch.id,
        calibrator_batch_no=scanned_batch.batch_no,
        calibrator_expiry_date=scanned_batch.expiry_date.isoformat() if scanned_batch.expiry_date else None,
        calibrator_available_min_quantity=_batch_total_min_units(scanned_batch),
        calibrator_min_unit=_clean_text(scanned_batch.min_unit) or "次",
        target_reagent_id=target.id if target else None,
        target_reagent_name=target.name if target else None,
        target_options=target_options,
        old_batch_candidates=batch_candidates,
        new_batch_candidates=batch_candidates,
        default_old_batch_no=default_old_batch_no,
        default_new_batch_no=default_new_batch_no,
        default_usage_times=1,
        result_unit=(_clean_text(target.calibration_result_unit) or "nmmol/L") if target else "nmmol/L",
        raw_barcode=barcode,
    )


@router.get("/calibration/mobile/preview", response_model=CalibrationScanPreviewResponse)
def mobile_calibration_preview(target_reagent_id: int = Query(..., ge=1), db: Session = Depends(get_db)):
    target = db.query(Reagent).filter(Reagent.id == target_reagent_id).first()
    if not target or target.is_archived or target.is_calibrator:
        raise HTTPException(status_code=404, detail="被校准试剂不存在")

    target_name = _normalize_name(target.name)
    calibrators = (
        db.query(Reagent)
        .filter(
            Reagent.instrument_id == target.instrument_id,
            Reagent.is_archived.is_(False),
            Reagent.is_calibrator.is_(True),
        )
        .order_by(Reagent.name.asc(), Reagent.id.asc())
        .all()
    )
    matched_source: Reagent | None = None
    matched_batch: BatchStock | None = None
    for candidate in calibrators:
        if _normalize_name(_strip_calibrator_suffix(candidate.name)) != target_name:
            continue
        available_batches = _load_available_batches(db, candidate.id)
        if not available_batches:
            continue
        matched_source = candidate
        matched_batch = available_batches[0]
        break

    calibration_source = "calibrator"
    if not matched_source or not matched_batch:
        available_target_batches = _load_available_batches(db, target.id, newest_first=True)
        if not available_target_batches:
            raise HTTPException(status_code=404, detail="当前试剂没有可用库存批号，不能进入校准流程")
        calibration_source = "reagent"
        matched_source = target
        matched_batch = available_target_batches[0]

    batch_candidates = _load_target_batch_candidates(db, target.id)
    if calibration_source == "reagent":
        default_old_batch_no, default_new_batch_no = _resolve_scanned_reagent_batch_values(
            batch_candidates,
            matched_batch.batch_no,
        )
    else:
        default_old_batch_no, default_new_batch_no = _resolve_default_batch_values(batch_candidates)

    return CalibrationScanPreviewResponse(
        calibration_source=calibration_source,
        calibrator_reagent_id=matched_source.id,
        calibrator_name=matched_source.name,
        calibrator_batch_stock_id=matched_batch.id,
        calibrator_batch_no=matched_batch.batch_no,
        calibrator_expiry_date=matched_batch.expiry_date.isoformat() if matched_batch.expiry_date else None,
        calibrator_available_min_quantity=_batch_total_min_units(matched_batch),
        calibrator_min_unit=_clean_text(matched_batch.min_unit) or "次",
        target_reagent_id=target.id,
        target_reagent_name=target.name,
        target_options=[_serialize_target_candidate(target)],
        old_batch_candidates=batch_candidates,
        new_batch_candidates=batch_candidates,
        default_old_batch_no=default_old_batch_no,
        default_new_batch_no=default_new_batch_no,
        default_usage_times=1,
        result_unit=_clean_text(target.calibration_result_unit) or "nmmol/L",
        raw_barcode=f"MOBILE::CALIBRATION::{target.id}",
    )


@router.get("/calibration/target-preview", response_model=CalibrationTargetPreviewResponse)
def calibration_target_preview(reagent_id: int = Query(..., ge=1), db: Session = Depends(get_db)):
    target = db.query(Reagent).filter(Reagent.id == reagent_id).first()
    if not target or target.is_archived or target.is_calibrator:
        raise HTTPException(status_code=404, detail="被校准试剂不存在")
    return _build_target_preview_payload(target, db)


@router.get("/calibration/records", response_model=list[CalibrationRecordItem])
def list_calibration_records(
    work_group_ids: Optional[str] = Query(None),
    instrument_ids: Optional[str] = Query(None),
    keyword: Optional[str] = Query(None),
    source: Optional[str] = Query(None),
    start_date: Optional[date] = Query(None),
    end_date: Optional[date] = Query(None),
    limit: int = Query(default=20, ge=1, le=100),
    db: Session = Depends(get_db),
):
    records = (
        _build_calibration_record_query(
            db,
            work_group_ids=work_group_ids,
            instrument_ids=instrument_ids,
            keyword=keyword,
            source=source,
            start_date=start_date,
            end_date=end_date,
        )
        .order_by(CalibrationRecord.created_at.desc(), CalibrationRecord.id.desc())
        .limit(limit)
        .all()
    )
    return [_serialize_record(record, db) for record in records]


@router.get("/calibration/records/export")
def export_calibration_records(
    work_group_ids: Optional[str] = Query(None),
    instrument_ids: Optional[str] = Query(None),
    keyword: Optional[str] = Query(None),
    source: Optional[str] = Query(None),
    start_date: Optional[date] = Query(None),
    end_date: Optional[date] = Query(None),
    db: Session = Depends(get_db),
):
    records = (
        _build_calibration_record_query(
            db,
            work_group_ids=work_group_ids,
            instrument_ids=instrument_ids,
            keyword=keyword,
            source=source,
            start_date=start_date,
            end_date=end_date,
        )
        .order_by(CalibrationRecord.created_at.desc(), CalibrationRecord.id.desc())
        .all()
    )
    buffer = _build_calibration_export_workbook(records, start_date, end_date)
    return _excel_response(buffer, _calibration_filename(start_date, end_date))


@router.get("/calibration/records/{record_id}", response_model=CalibrationRecordItem)
def get_calibration_record(record_id: int, db: Session = Depends(get_db)):
    record = db.query(CalibrationRecord).filter(CalibrationRecord.id == record_id).first()
    if not record:
        raise HTTPException(status_code=404, detail="校准记录不存在")
    return _serialize_record(record, db)


@router.post("/calibration/records", response_model=CalibrationRecordItem, status_code=201)
def create_calibration_record(data: CalibrationRecordCreate, db: Session = Depends(get_db)):
    operator = _require_operator(data.operator, db)
    calibrator = db.query(Reagent).filter(Reagent.id == data.calibrator_reagent_id).first()
    if not calibrator or calibrator.is_archived:
        raise HTTPException(status_code=400, detail="校准来源试剂不存在或已归档")

    target = db.query(Reagent).filter(Reagent.id == data.target_reagent_id).first()
    if not target or target.is_archived or target.is_calibrator:
        raise HTTPException(status_code=400, detail="被校准试剂不存在，或当前选择的是校准品")

    if not calibrator.is_calibrator and calibrator.id != target.id:
        raise HTTPException(status_code=400, detail="普通试剂扫码校准时，被校准试剂必须是当前扫码试剂")

    calibrator_batch = (
        db.query(BatchStock)
        .filter(
            BatchStock.id == data.calibrator_batch_stock_id,
            BatchStock.is_archived.is_(False),
        )
        .first()
    )
    if not calibrator_batch or calibrator_batch.reagent_id != calibrator.id:
        raise HTTPException(status_code=400, detail="校准来源批号不存在或与扫码试剂不匹配")

    should_deduct_calibration_source = calibrator.is_calibrator
    if should_deduct_calibration_source:
        available_min_quantity = _batch_total_min_units(calibrator_batch)
        if available_min_quantity < data.usage_times:
            raise HTTPException(status_code=400, detail="校准品库存不足，不能完成本次校准")

        _set_batch_total_min_units(calibrator_batch, available_min_quantity - data.usage_times)

    transaction = Transaction(
        reagent_id=calibrator.id,
        batch_stock_id=calibrator_batch.id,
        type="CALIBRATION",
        quantity=data.usage_times,
        operator=operator,
        source="scan",
        raw_barcode=_clean_text(data.raw_barcode) or f"CAL::{calibrator.id}",
        package_unit=calibrator_batch.package_unit,
        min_unit=calibrator_batch.min_unit,
        package_to_min_qty=_safe_ratio(calibrator_batch.package_to_min_qty),
        remark=_calibration_transaction_remark(
            target_reagent_id=target.id,
            target_reagent_name=target.name,
            target_old_batch_no=data.target_old_batch_no,
            target_new_batch_no=data.target_new_batch_no,
            usage_times=data.usage_times,
            result_unit=_clean_text(data.result_unit) or _clean_text(target.calibration_result_unit) or "nmmol/L",
        ),
    )
    db.add(transaction)
    db.flush()

    record = CalibrationRecord(
        calibrator_reagent_id=calibrator.id,
        calibrator_batch_stock_id=calibrator_batch.id,
        target_reagent_id=target.id,
        transaction_id=transaction.id,
        raw_barcode=_clean_text(data.raw_barcode) or None,
        target_old_batch_no=_clean_text(data.target_old_batch_no) or None,
        target_new_batch_no=_clean_text(data.target_new_batch_no) or None,
        usage_times=data.usage_times,
        result_unit=_clean_text(data.result_unit) or _clean_text(target.calibration_result_unit) or "nmmol/L",
        operator=operator,
        old_result_1=_clean_text(data.old_result_1) or None,
        old_result_2=_clean_text(data.old_result_2) or None,
        old_result_3=_clean_text(data.old_result_3) or None,
        old_result_4=_clean_text(data.old_result_4) or None,
        old_result_5=_clean_text(data.old_result_5) or None,
        new_result_1=_clean_text(data.new_result_1) or None,
        new_result_2=_clean_text(data.new_result_2) or None,
        new_result_3=_clean_text(data.new_result_3) or None,
        new_result_4=_clean_text(data.new_result_4) or None,
        new_result_5=_clean_text(data.new_result_5) or None,
    )
    db.add(record)
    db.flush()

    if data.auto_outbound_package:
        _create_linked_outbound_transaction(
            source_reagent=calibrator,
            source_batch=calibrator_batch,
            operator=operator,
            raw_barcode=data.raw_barcode,
            calibration_record_id=record.id,
            db=db,
            outbound_reagent_id=target.id,
            outbound_batch_stock_id=data.outbound_batch_stock_id,
            outbound_quantity=data.outbound_quantity,
            outbound_unit=data.outbound_unit,
        )

    db.commit()
    db.refresh(record)
    return _serialize_record(record, db)


@router.put("/calibration/records/{record_id}", response_model=CalibrationRecordItem)
def update_calibration_record(record_id: int, data: CalibrationRecordUpdate, db: Session = Depends(get_db)):
    record = db.query(CalibrationRecord).filter(CalibrationRecord.id == record_id).first()
    if not record:
        raise HTTPException(status_code=404, detail="校准记录不存在")

    operator = _require_operator(data.operator, db)
    calibrator = record.calibrator_reagent
    target = db.query(Reagent).filter(Reagent.id == data.target_reagent_id).first()
    if not target or target.is_archived or target.is_calibrator:
        raise HTTPException(status_code=400, detail="被校准试剂不存在，或当前选择的是校准品")
    if calibrator and not calibrator.is_calibrator and calibrator.id != target.id:
        raise HTTPException(status_code=400, detail="普通试剂扫码校准时，被校准试剂必须是当前扫码试剂")

    calibrator_batch = db.query(BatchStock).filter(BatchStock.id == record.calibrator_batch_stock_id).first()
    if not calibrator_batch:
        raise HTTPException(status_code=400, detail="校准品批次不存在，无法编辑本次校准")

    delta_usage = data.usage_times - record.usage_times
    if calibrator and calibrator.is_calibrator:
        if delta_usage > 0:
            available_min_quantity = _batch_total_min_units(calibrator_batch)
            if available_min_quantity < delta_usage:
                raise HTTPException(status_code=400, detail="校准品库存不足，不能增加本次校准用量")
            _set_batch_total_min_units(calibrator_batch, available_min_quantity - delta_usage)
        elif delta_usage < 0:
            _set_batch_total_min_units(calibrator_batch, _batch_total_min_units(calibrator_batch) + abs(delta_usage))

    record.target_reagent_id = target.id
    record.target_old_batch_no = _clean_text(data.target_old_batch_no) or None
    record.target_new_batch_no = _clean_text(data.target_new_batch_no) or None
    record.usage_times = data.usage_times
    record.result_unit = _clean_text(data.result_unit) or _clean_text(target.calibration_result_unit) or "nmmol/L"
    record.operator = operator
    record.old_result_1 = _clean_text(data.old_result_1) or None
    record.old_result_2 = _clean_text(data.old_result_2) or None
    record.old_result_3 = _clean_text(data.old_result_3) or None
    record.old_result_4 = _clean_text(data.old_result_4) or None
    record.old_result_5 = _clean_text(data.old_result_5) or None
    record.new_result_1 = _clean_text(data.new_result_1) or None
    record.new_result_2 = _clean_text(data.new_result_2) or None
    record.new_result_3 = _clean_text(data.new_result_3) or None
    record.new_result_4 = _clean_text(data.new_result_4) or None
    record.new_result_5 = _clean_text(data.new_result_5) or None

    if record.transaction_id:
        transaction = db.query(Transaction).filter(Transaction.id == record.transaction_id).first()
        if transaction:
            transaction.quantity = data.usage_times
            transaction.operator = operator
            transaction.remark = _calibration_transaction_remark(
                target_reagent_id=target.id,
                target_reagent_name=target.name,
                target_old_batch_no=record.target_old_batch_no,
                target_new_batch_no=record.target_new_batch_no,
                usage_times=data.usage_times,
                result_unit=record.result_unit,
            )

    db.commit()
    db.refresh(record)
    return _serialize_record(record, db)
