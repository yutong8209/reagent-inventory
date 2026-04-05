from datetime import date, datetime
from io import BytesIO
from urllib.parse import unquote

import models  # noqa: F401
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from api import reports
from db import get_db
from db.base import Base
from models import BatchStock, Instrument, Reagent, Transaction, WorkGroup
from services import ensure_default_catalog, ensure_default_operators


def build_client(tmp_path):
    db_path = tmp_path / "test.db"
    engine = create_engine(f"sqlite:///{db_path}", connect_args={"check_same_thread": False})
    TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    Base.metadata.create_all(bind=engine)

    with TestingSessionLocal() as db:
        ensure_default_catalog(db)
        ensure_default_operators(db)

    app = FastAPI()
    app.include_router(reports.router, prefix="/api")

    def override_get_db():
        db = TestingSessionLocal()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db
    return TestClient(app), TestingSessionLocal


def seed_report_data(SessionLocal):
    with SessionLocal() as db:
        work_groups = db.query(WorkGroup).order_by(WorkGroup.id.asc()).all()
        primary_group = work_groups[0]
        secondary_group = work_groups[1]

        primary_instrument = (
            db.query(Instrument)
            .filter(Instrument.work_group_id == primary_group.id)
            .order_by(Instrument.id.asc())
            .first()
        )
        secondary_instrument = (
            db.query(Instrument)
            .filter(Instrument.work_group_id == secondary_group.id)
            .order_by(Instrument.id.asc())
            .first()
        )

        active_reagent = Reagent(
            instrument_id=primary_instrument.id,
            name="甲功试剂A",
            match_mode="exact",
            match_pattern="REPORT-001",
            raw_barcode_sample="REPORT-001",
            manufacturer="厂家甲",
            unit="瓶",
            package_unit="盒",
            package_to_min_qty=2,
            warning_days=30,
            warning_stock=5,
            follow_instrument=False,
            is_archived=False,
        )
        archived_reagent = Reagent(
            instrument_id=secondary_instrument.id,
            name="归档试剂B",
            match_mode="exact",
            match_pattern="REPORT-002",
            raw_barcode_sample="REPORT-002",
            manufacturer="厂家乙",
            unit="支",
            package_unit="盒",
            package_to_min_qty=1,
            warning_days=30,
            warning_stock=5,
            follow_instrument=False,
            is_archived=True,
        )
        db.add_all([active_reagent, archived_reagent])
        db.flush()

        active_batch = BatchStock(
            reagent_id=active_reagent.id,
            batch_no="ACTIVE-BATCH",
            expiry_date=date(2031, 12, 31),
            quantity=10,
            loose_min_quantity=0,
            package_unit="盒",
            min_unit="瓶",
            package_to_min_qty=2,
        )
        archived_batch = BatchStock(
            reagent_id=archived_reagent.id,
            batch_no="ARCHIVED-BATCH",
            expiry_date=date(2031, 12, 31),
            quantity=5,
            loose_min_quantity=0,
            package_unit="盒",
            min_unit="支",
            package_to_min_qty=1,
        )
        db.add_all([active_batch, archived_batch])
        db.flush()

        db.add_all([
            Transaction(
                reagent_id=active_reagent.id,
                batch_stock_id=active_batch.id,
                type="IN",
                quantity=3,
                operator="陈茹",
                source="manual",
                raw_barcode="REPORT-001",
                package_unit="盒",
                min_unit="瓶",
                package_to_min_qty=2,
                remark=None,
                created_at=datetime(2026, 3, 10, 9, 30, 0),
            ),
            Transaction(
                reagent_id=active_reagent.id,
                batch_stock_id=active_batch.id,
                type="OUT",
                quantity=4,
                operator="陈茹",
                source="scan",
                raw_barcode="REPORT-001",
                package_unit="盒",
                min_unit="瓶",
                package_to_min_qty=2,
                remark='{"deducted_min_quantity":4}',
                created_at=datetime(2026, 3, 15, 15, 20, 0),
            ),
            Transaction(
                reagent_id=archived_reagent.id,
                batch_stock_id=archived_batch.id,
                type="OUT",
                quantity=7,
                operator="陈茹",
                source="manual",
                raw_barcode="REPORT-002",
                package_unit="盒",
                min_unit="支",
                package_to_min_qty=1,
                remark='{"deducted_min_quantity":7}',
                created_at=datetime(2026, 3, 20, 18, 45, 0),
            ),
        ])
        db.commit()

        return {
            "primary_group_id": primary_group.id,
            "primary_instrument_id": primary_instrument.id,
            "active_reagent_id": active_reagent.id,
            "archived_reagent_id": archived_reagent.id,
        }


def test_transactions_query_and_export_support_filters(tmp_path):
    client, SessionLocal = build_client(tmp_path)
    ids = seed_report_data(SessionLocal)

    response = client.get(
        "/api/transactions",
        params={
            "work_group_id": ids["primary_group_id"],
            "instrument_id": ids["primary_instrument_id"],
            "reagent_id": ids["active_reagent_id"],
            "keyword": "甲功",
            "type": "OUT",
            "start_date": "2026-03-01",
            "end_date": "2026-03-31",
            "limit": 20,
        },
    )
    assert response.status_code == 200
    payload = response.json()
    assert len(payload) == 1
    assert payload[0]["reagent_name"] == "甲功试剂A"
    assert payload[0]["quantity_min"] == 4
    assert payload[0]["batch_no"] == "ACTIVE-BATCH"

    export_response = client.get(
        "/api/transactions/export",
        params={
            "work_group_id": ids["primary_group_id"],
            "start_date": "2026-03-01",
            "end_date": "2026-03-31",
        },
    )
    assert export_response.status_code == 200
    assert "流水_2026-03-01_2026-03-31.xlsx" in unquote(export_response.headers["content-disposition"])

    workbook = load_workbook(BytesIO(export_response.content))
    sheet = workbook.active
    assert sheet.cell(1, 1).value == "流水时间"
    assert sheet.max_row >= 3


def test_usage_query_and_export_include_archived_history(tmp_path):
    client, SessionLocal = build_client(tmp_path)
    ids = seed_report_data(SessionLocal)

    response = client.get(
        "/api/usage",
        params={
            "keyword": "试剂",
            "start_date": "2026-03-01",
            "end_date": "2026-03-31",
        },
    )
    assert response.status_code == 200
    payload = response.json()
    names = [item["reagent_name"] for item in payload]
    assert "甲功试剂A" in names
    assert "归档试剂B" in names

    archived_item = next(item for item in payload if item["reagent_id"] == ids["archived_reagent_id"])
    assert archived_item["total_usage_min"] == 7
    assert archived_item["min_unit"] == "支"

    export_response = client.get(
        "/api/usage/export",
        params={
            "start_date": "2026-03-01",
            "end_date": "2026-03-31",
        },
    )
    assert export_response.status_code == 200
    assert "用量_2026-03-01_2026-03-31.xlsx" in unquote(export_response.headers["content-disposition"])

    workbook = load_workbook(BytesIO(export_response.content))
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][0] == "工作组"
    assert any(row[2] == "归档试剂B" for row in rows[1:])


def test_reporting_reagent_options_include_archived_records(tmp_path):
    client, SessionLocal = build_client(tmp_path)
    ids = seed_report_data(SessionLocal)

    response = client.get("/api/reporting/reagents", params={"keyword": "试剂"})
    assert response.status_code == 200
    payload = response.json()
    archived_item = next(item for item in payload if item["id"] == ids["archived_reagent_id"])
    assert archived_item["is_archived"] is True
    assert archived_item["name"] == "归档试剂B"
