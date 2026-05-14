from io import BytesIO
from datetime import date, datetime

import models  # noqa: F401
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from api import barcode, inbound, inventory
from db import get_db
from db.base import Base
from models import BarcodeBatchRule, BatchStock, CalibrationRecord, Instrument, Operator, ParseLog, Reagent, Transaction, WorkGroup
from services import ensure_default_catalog, ensure_default_operators, ensure_default_unit_options


def build_client(tmp_path):
    db_path = tmp_path / "test.db"
    engine = create_engine(f"sqlite:///{db_path}", connect_args={"check_same_thread": False})
    TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    Base.metadata.create_all(bind=engine)

    with TestingSessionLocal() as db:
        ensure_default_catalog(db)
        ensure_default_operators(db)
        ensure_default_unit_options(db)

    app = FastAPI()
    app.include_router(barcode.router, prefix="/api")
    app.include_router(inbound.router, prefix="/api")
    app.include_router(inventory.router, prefix="/api")

    def override_get_db():
        db = TestingSessionLocal()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db
    return TestClient(app), TestingSessionLocal


def get_first_catalog_session(SessionLocal):
    with SessionLocal() as db:
        work_group = db.query(WorkGroup).order_by(WorkGroup.id.asc()).first()
        instrument = (
            db.query(Instrument)
            .filter(Instrument.work_group_id == work_group.id)
            .order_by(Instrument.id.asc())
            .first()
        )
        return work_group.id, instrument.id


def test_reagent_list_filters_and_keyword(tmp_path):
    client, SessionLocal = build_client(tmp_path)
    work_group_id, instrument_id = get_first_catalog_session(SessionLocal)

    create_response = client.post("/api/reagents", json={
        "instrument_id": instrument_id,
        "name": "丙氨酸氨基转移酶试剂",
        "match_mode": "prefix",
        "match_pattern": "01999911112222",
        "manufacturer": "测试厂家",
        "unit": "瓶",
        "package_unit": "盒",
        "package_to_min_qty": 2,
        "warning_days": 20,
        "warning_stock": 3,
    })
    assert create_response.status_code == 201

    list_response = client.get(f"/api/reagents?work_group_id={work_group_id}&instrument_id={instrument_id}&keyword=丙氨酸")
    assert list_response.status_code == 200
    payload = list_response.json()
    assert len(payload) == 1
    assert payload[0]["name"] == "丙氨酸氨基转移酶试剂"
    assert payload[0]["manufacturer"] == "测试厂家"


def test_reagent_group_crud_and_reagent_assignment(tmp_path):
    client, SessionLocal = build_client(tmp_path)
    _, instrument_id = get_first_catalog_session(SessionLocal)

    create_group_response = client.post("/api/reagent-groups", json={
        "name": "甲功",
        "sort_order": 10,
        "is_active": True,
    })
    assert create_group_response.status_code == 201
    group_one_id = create_group_response.json()["id"]

    create_second_group_response = client.post("/api/reagent-groups", json={
        "name": "激素",
        "sort_order": 20,
        "is_active": True,
    })
    assert create_second_group_response.status_code == 201
    group_two_id = create_second_group_response.json()["id"]

    create_reagent_response = client.post("/api/reagents", json={
        "instrument_id": instrument_id,
        "name": "项目分组试剂",
        "match_mode": "exact",
        "match_pattern": "GROUP-REAGENT-001",
        "manufacturer": "分组厂家",
        "unit": "瓶",
        "package_unit": "盒",
        "package_to_min_qty": 2,
        "warning_days": 30,
        "warning_stock": 5,
        "reagent_group_id": group_one_id,
    })
    assert create_reagent_response.status_code == 201
    reagent_id = create_reagent_response.json()["id"]
    assert create_reagent_response.json()["reagent_group_id"] == group_one_id

    list_response = client.get(f"/api/reagents?instrument_id={instrument_id}&keyword=项目分组")
    assert list_response.status_code == 200
    assert list_response.json()[0]["reagent_group_id"] == group_one_id
    assert list_response.json()[0]["reagent_group_name"] == "甲功"

    update_reagent_response = client.put(f"/api/reagents/{reagent_id}", json={
        "instrument_id": instrument_id,
        "name": "项目分组试剂",
        "manufacturer": "分组厂家",
        "package_unit": "盒",
        "unit": "瓶",
        "package_to_min_qty": 2,
        "warning_days": 30,
        "warning_stock": 5,
        "follow_instrument": False,
        "is_calibrator": False,
        "calibration_result_unit": "nmmol/L",
        "reagent_group_id": group_two_id,
    })
    assert update_reagent_response.status_code == 200
    assert update_reagent_response.json()["reagent_group_id"] == group_two_id
    assert update_reagent_response.json()["reagent_group_name"] == "激素"

    deactivate_group_response = client.put(f"/api/reagent-groups/{group_one_id}", json={
        "name": "甲功",
        "sort_order": 10,
        "is_active": False,
    })
    assert deactivate_group_response.status_code == 200
    assert deactivate_group_response.json()["is_active"] is False

    active_groups_response = client.get("/api/reagent-groups")
    assert active_groups_response.status_code == 200
    active_group_names = {item["name"] for item in active_groups_response.json()}
    assert "甲功" not in active_group_names
    assert "激素" in active_group_names

    all_groups_response = client.get("/api/reagent-groups?include_inactive=true")
    assert all_groups_response.status_code == 200
    all_groups = {item["name"]: item for item in all_groups_response.json()}
    assert all_groups["甲功"]["is_active"] is False
    assert all_groups["激素"]["is_active"] is True


def test_operator_crud(tmp_path):
    client, _ = build_client(tmp_path)

    list_response = client.get("/api/operators")
    assert list_response.status_code == 200
    assert len(list_response.json()) >= 1

    create_response = client.post("/api/operators", json={"name": "测试操作员"})
    assert create_response.status_code == 201
    operator_id = create_response.json()["id"]

    update_response = client.put(f"/api/operators/{operator_id}", json={
        "name": "测试操作员A",
        "is_active": False,
    })
    assert update_response.status_code == 200
    assert update_response.json()["name"] == "测试操作员A"
    assert update_response.json()["is_active"] is False

    delete_response = client.delete(f"/api/operators/{operator_id}")
    assert delete_response.status_code == 200
    assert delete_response.json()["success"] is True
    assert delete_response.json()["operator_id"] == operator_id

    list_after_delete = client.get("/api/operators")
    assert list_after_delete.status_code == 200
    assert all(item["id"] != operator_id for item in list_after_delete.json())


def test_first_manual_inbound_creates_reagent_and_batch(tmp_path):
    client, SessionLocal = build_client(tmp_path)
    work_group_id, instrument_id = get_first_catalog_session(SessionLocal)

    response = client.post("/api/inbound/batch", json={
        "operator": "陈茜",
        "items": [{
            "source": "manual",
            "work_group_id": work_group_id,
            "instrument_id": instrument_id,
            "reagent_name": "首次手工建档试剂",
            "manufacturer": "手工厂家",
            "batch_no": "BATCH-001",
            "expiry_date": "2030-12-31",
            "quantity": 4,
            "package_unit": "盒",
            "min_unit": "瓶",
            "package_to_min_qty": 6,
            "warning_days": 12,
            "warning_stock": 2,
            "raw_barcode": "MANUAL::TEST::001"
        }]
    })
    assert response.status_code == 200
    assert response.json()["item_count"] == 1

    with SessionLocal() as db:
        reagent = db.query(Reagent).filter(Reagent.name == "首次手工建档试剂").first()
        assert reagent is not None
        assert reagent.manufacturer == "手工厂家"
        assert reagent.package_to_min_qty == 6
        assert reagent.warning_days == 12
        assert reagent.warning_stock == 2
        batch = db.query(BatchStock).filter(BatchStock.reagent_id == reagent.id).first()
        assert batch is not None
        assert batch.quantity == 4


def test_existing_reagent_inbound_reuses_archive_package_info(tmp_path):
    client, SessionLocal = build_client(tmp_path)
    _, instrument_id = get_first_catalog_session(SessionLocal)

    create_response = client.post("/api/reagents", json={
        "instrument_id": instrument_id,
        "name": "已有试剂",
        "match_mode": "prefix",
        "match_pattern": "01888877776666",
        "manufacturer": "归档厂家",
        "unit": "支",
        "package_unit": "袋",
        "package_to_min_qty": 8,
        "warning_days": 15,
        "warning_stock": 4,
    })
    reagent_id = create_response.json()["id"]

    response = client.post("/api/inbound/batch", json={
        "operator": "陈茜",
        "items": [{
            "source": "manual",
            "reagent_id": reagent_id,
            "reagent_name": "已有试剂",
            "batch_no": "KEEP-001",
            "expiry_date": "2031-01-01",
            "quantity": 2,
            "package_unit": "箱",
            "min_unit": "片",
            "package_to_min_qty": 99,
            "raw_barcode": "MANUAL::KEEP::001"
        }]
    })
    assert response.status_code == 200
    item = response.json()["items"][0]
    assert item["package_unit"] == "袋"
    assert item["min_unit"] == "支"
    assert item["package_to_min_qty"] == 8


def test_unknown_barcode_infers_instrument_from_product_family(tmp_path):
    client, SessionLocal = build_client(tmp_path)

    with SessionLocal() as db:
        work_group = WorkGroup(name="条码推断组")
        dxi = Instrument(name="DXI800化学发光仪", model="DXI800", work_group=work_group)
        maglumi = Instrument(name="MAGLUMI X8", model="X8", work_group=work_group)
        db.add_all([work_group, dxi, maglumi])
        db.flush()

        db.add_all([
            Reagent(
                instrument_id=dxi.id,
                name="TT4",
                match_mode="prefix",
                match_pattern="0115099590225742",
                raw_barcode_sample="0115099590225742112507311726073110538759",
                unit="瓶",
                package_unit="盒",
                package_to_min_qty=1,
                warning_days=30,
                warning_stock=5,
            ),
            Reagent(
                instrument_id=dxi.id,
                name="FER",
                match_mode="prefix",
                match_pattern="0115099590575052",
                raw_barcode_sample="0115099590575052112506301726063010538639",
                unit="瓶",
                package_unit="盒",
                package_to_min_qty=1,
                warning_days=30,
                warning_stock=5,
            ),
            Reagent(
                instrument_id=maglumi.id,
                name="HBsAg",
                match_mode="prefix",
                match_pattern="0106947145508814",
                raw_barcode_sample="0106947145508814112511201727052010123250711240130210009M",
                unit="瓶",
                package_unit="盒",
                package_to_min_qty=1,
                warning_days=30,
                warning_stock=5,
            ),
        ])
        db.commit()
        dxi_id = dxi.id
        work_group_id = work_group.id

    response = client.post(
        "/api/barcode/parse",
        json={"barcode": "0115099590227210112509301726093010538990"},
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["match_type"] == "none"
    assert payload["reagent_id"] is None
    assert payload["work_group_id"] == work_group_id
    assert payload["instrument_id"] == dxi_id
    assert payload["extracted_batch_no"] == "538990"
    assert payload["extracted_expiry_date"] == "2026-09-30"


def test_dxi_ft4_calibrator_barcode_uses_lot_after_ai_10(tmp_path):
    client, SessionLocal = build_client(tmp_path)

    barcode = "0115099590225810112507311726073110538615"
    with SessionLocal() as db:
        work_group = WorkGroup(name="DXI")
        instrument = Instrument(name="DXI800", model="DXI800", work_group=work_group)
        db.add_all([work_group, instrument])
        db.flush()
        db.add(
            Reagent(
                instrument_id=instrument.id,
                name="FT4 CAL",
                match_mode="prefix",
                match_pattern="0115099590225810",
                raw_barcode_sample=barcode,
                unit="次",
                package_unit="盒",
                package_to_min_qty=10,
                warning_days=30,
                warning_stock=5,
                is_calibrator=True,
            )
        )
        db.commit()

    response = client.post("/api/barcode/parse", json={"barcode": barcode})
    assert response.status_code == 200
    payload = response.json()
    assert payload["reagent_name"] == "FT4 CAL"
    assert payload["extracted_production_date"] == "2025-07-31"
    assert payload["extracted_expiry_date"] == "2026-07-31"
    assert payload["extracted_batch_no"] == "538615"
    assert payload["extracted_batch_no"] != "112507"


def test_reset_reagent_data_keeps_catalog_and_operators(tmp_path):
    client, SessionLocal = build_client(tmp_path)
    work_group_id, instrument_id = get_first_catalog_session(SessionLocal)

    with SessionLocal() as db:
        reagent = Reagent(
            instrument_id=instrument_id,
            name="待清空试剂",
            match_mode="exact",
            match_pattern="RESET-001",
            raw_barcode_sample="RESET-001",
            manufacturer="厂家",
            unit="瓶",
            package_unit="盒",
            package_to_min_qty=1,
            warning_days=30,
            warning_stock=5,
        )
        db.add(reagent)
        db.flush()

        batch = BatchStock(
            reagent_id=reagent.id,
            batch_no="RESET-BATCH",
            expiry_date=date(2030, 1, 1),
            quantity=1,
            package_unit="盒",
            min_unit="瓶",
            package_to_min_qty=1,
        )
        db.add(batch)
        db.flush()

        db.add(Transaction(
            reagent_id=reagent.id,
            batch_stock_id=batch.id,
            type="IN",
            quantity=1,
            operator="陈茜",
            source="manual",
            raw_barcode="RESET-001",
            package_unit="盒",
            min_unit="瓶",
            package_to_min_qty=1,
            remark=None,
        ))
        db.add(ParseLog(raw_barcode="RESET-001", match_result="exact", matched_reagent_id=reagent.id))
        db.commit()

    reset_response = client.post("/api/admin/reset-reagent-data", json={"password": "wtsyyjyk"})
    assert reset_response.status_code == 200
    assert reset_response.json()["success"] is True

    with SessionLocal() as db:
        assert db.query(WorkGroup).filter(WorkGroup.id == work_group_id).count() == 1
        assert db.query(Instrument).filter(Instrument.id == instrument_id).count() == 1
        assert db.query(Operator).count() >= 1
        assert db.query(Reagent).count() == 0
        assert db.query(BatchStock).count() == 0
        assert db.query(Transaction).count() == 0
        assert db.query(ParseLog).count() == 0


def test_inventory_uses_instrument_defaults_until_reagent_override(tmp_path):
    client, SessionLocal = build_client(tmp_path)
    _, instrument_id = get_first_catalog_session(SessionLocal)

    with SessionLocal() as db:
        instrument = db.query(Instrument).filter(Instrument.id == instrument_id).first()
        instrument.default_warning_days = 20
        instrument.default_warning_stock = 10

        reagent = Reagent(
            instrument_id=instrument_id,
            name="预警继承试剂",
            match_mode="exact",
            match_pattern="WARN-INHERIT-001",
            raw_barcode_sample="WARN-INHERIT-001",
            manufacturer="测试厂家",
            unit="瓶",
            package_unit="盒",
            package_to_min_qty=1,
            warning_days=2,
            warning_stock=3,
            follow_instrument=True,
        )
        db.add(reagent)
        db.flush()

        db.add(BatchStock(
            reagent_id=reagent.id,
            batch_no="WARN-BATCH-1",
            expiry_date=date.today().replace(year=date.today().year + 1) if date.today().month != 2 or date.today().day != 29 else date(date.today().year + 1, 2, 28),
            quantity=8,
            package_unit="盒",
            min_unit="瓶",
            package_to_min_qty=1,
        ))
        db.commit()
        reagent_id = reagent.id

    with SessionLocal() as db:
        batch = db.query(BatchStock).filter(BatchStock.reagent_id == reagent_id).first()
        batch.expiry_date = date.today().fromordinal(date.today().toordinal() + 5)
        db.commit()

    inventory_response = client.get("/api/inventory")
    assert inventory_response.status_code == 200
    payload = inventory_response.json()
    reagent_payload = next(
        reagent
        for group in payload
        for instrument in group["instruments"]
        for reagent in instrument["reagents"]
        if reagent["id"] == reagent_id
    )
    assert reagent_payload["warning_source"] == "instrument"
    assert reagent_payload["effective_warning_days"] == 20
    assert reagent_payload["effective_warning_stock"] == 10
    assert reagent_payload["single_batch_warning_stock"] is None
    assert reagent_payload["effective_single_batch_warning_stock"] == 10
    assert reagent_payload["status_level"] == "expiring"

    update_response = client.put(f"/api/reagents/{reagent_id}/warning-settings", json={
        "follow_instrument": False,
        "warning_days": 2,
        "warning_stock": 5,
        "single_batch_warning_stock": 3,
    })
    assert update_response.status_code == 200
    assert update_response.json()["warning_source"] == "reagent"
    assert update_response.json()["single_batch_warning_stock"] is None
    assert update_response.json()["effective_single_batch_warning_stock"] == 5

    inventory_response = client.get("/api/inventory")
    reagent_payload = next(
        reagent
        for group in inventory_response.json()
        for instrument in group["instruments"]
        for reagent in instrument["reagents"]
        if reagent["id"] == reagent_id
    )
    assert reagent_payload["warning_source"] == "reagent"
    assert reagent_payload["effective_warning_days"] == 2
    assert reagent_payload["effective_warning_stock"] == 5
    assert reagent_payload["single_batch_warning_stock"] is None
    assert reagent_payload["effective_single_batch_warning_stock"] == 5
    assert reagent_payload["status_level"] == "normal"


def test_apply_instrument_warning_defaults_switches_reagents_back_to_follow(tmp_path):
    client, SessionLocal = build_client(tmp_path)
    _, instrument_id = get_first_catalog_session(SessionLocal)

    with SessionLocal() as db:
        instrument = db.query(Instrument).filter(Instrument.id == instrument_id).first()
        instrument.default_warning_days = 6
        instrument.default_warning_stock = 7

        reagent = Reagent(
            instrument_id=instrument_id,
            name="批量切回跟随试剂",
            match_mode="exact",
            match_pattern="WARN-APPLY-001",
            raw_barcode_sample="WARN-APPLY-001",
            manufacturer="测试厂家",
            unit="瓶",
            package_unit="盒",
            package_to_min_qty=1,
            warning_days=1,
            warning_stock=2,
            follow_instrument=False,
        )
        db.add(reagent)
        db.flush()

        db.add_all([
            BatchStock(
                reagent_id=reagent.id,
                batch_no="APPLY-BATCH-2",
                expiry_date=date.today().fromordinal(date.today().toordinal() + 9),
                quantity=2,
                package_unit="盒",
                min_unit="瓶",
                package_to_min_qty=1,
            ),
            BatchStock(
                reagent_id=reagent.id,
                batch_no="APPLY-BATCH-1",
                expiry_date=date.today().fromordinal(date.today().toordinal() + 5),
                quantity=3,
                package_unit="盒",
                min_unit="瓶",
                package_to_min_qty=1,
            ),
        ])
        db.commit()
        reagent_id = reagent.id

    defaults_response = client.put(f"/api/instruments/{instrument_id}/warning-defaults", json={
        "default_warning_days": 10,
        "default_warning_stock": 4,
    })
    assert defaults_response.status_code == 200

    apply_response = client.post(f"/api/instruments/{instrument_id}/warning-defaults/apply")
    assert apply_response.status_code == 200
    assert apply_response.json()["success"] is True
    assert apply_response.json()["updated_count"] >= 1

    inventory_response = client.get("/api/inventory")
    reagent_payload = next(
        reagent
        for group in inventory_response.json()
        for instrument in group["instruments"]
        for reagent in instrument["reagents"]
        if reagent["id"] == reagent_id
    )
    assert reagent_payload["warning_source"] == "instrument"
    assert reagent_payload["follow_instrument"] is True
    assert reagent_payload["effective_warning_days"] == 10
    assert reagent_payload["status_level"] == "expiring"
    assert [batch["batch_no"] for batch in reagent_payload["batches"]] == ["APPLY-BATCH-1", "APPLY-BATCH-2"]


def test_batch_update_instrument_warnings_updates_existing_reagents_and_defaults(tmp_path):
    client, SessionLocal = build_client(tmp_path)
    _, instrument_id = get_first_catalog_session(SessionLocal)

    with SessionLocal() as db:
        reagent = Reagent(
            instrument_id=instrument_id,
            name="批量预警试剂",
            match_mode="exact",
            match_pattern="WARN-BULK-001",
            raw_barcode_sample="WARN-BULK-001",
            manufacturer="测试厂家",
            unit="瓶",
            package_unit="盒",
            package_to_min_qty=1,
            warning_days=3,
            warning_stock=2,
            follow_instrument=True,
        )
        db.add(reagent)
        db.commit()
        reagent_id = reagent.id

    response = client.post(f"/api/instruments/{instrument_id}/warning-batch-update", json={
        "default_warning_days": 18,
        "default_warning_stock": 9,
    })
    assert response.status_code == 200
    payload = response.json()
    assert payload["success"] is True
    assert payload["updated_count"] >= 1

    with SessionLocal() as db:
        instrument = db.query(Instrument).filter(Instrument.id == instrument_id).first()
        reagent = db.query(Reagent).filter(Reagent.id == reagent_id).first()
        assert instrument.default_warning_days == 18
        assert instrument.default_warning_stock == 9
        assert reagent.warning_days == 18
        assert reagent.warning_stock == 9
        assert reagent.follow_instrument is False


def test_dashboard_overview_returns_pending_tasks_warnings_and_recent_outbounds(tmp_path):
    client, SessionLocal = build_client(tmp_path)
    _, instrument_id = get_first_catalog_session(SessionLocal)

    with SessionLocal() as db:
        reagent_expired = Reagent(
            instrument_id=instrument_id,
            name="DASH-EXPIRED",
            match_mode="exact",
            match_pattern="DASH-EXPIRED",
            raw_barcode_sample="DASH-EXPIRED",
            manufacturer="Maker A",
            unit="瓶",
            package_unit="盒",
            package_to_min_qty=1,
            warning_days=30,
            warning_stock=5,
            follow_instrument=False,
        )
        reagent_empty = Reagent(
            instrument_id=instrument_id,
            name="DASH-EMPTY",
            match_mode="exact",
            match_pattern="DASH-EMPTY",
            raw_barcode_sample="DASH-EMPTY",
            manufacturer="Maker B",
            unit="瓶",
            package_unit="盒",
            package_to_min_qty=1,
            warning_days=30,
            warning_stock=5,
            follow_instrument=False,
        )
        reagent_normal = Reagent(
            instrument_id=instrument_id,
            name="DASH-NORMAL",
            match_mode="exact",
            match_pattern="DASH-NORMAL",
            raw_barcode_sample="DASH-NORMAL",
            manufacturer="Maker C",
            unit="瓶",
            package_unit="盒",
            package_to_min_qty=1,
            warning_days=5,
            warning_stock=1,
            follow_instrument=False,
        )
        reagent_zero_batch_calibrator = Reagent(
            instrument_id=instrument_id,
            name="DASH-ZERO-BATCH-CAL",
            match_mode="exact",
            match_pattern="DASH-ZERO-BATCH-CAL",
            raw_barcode_sample="DASH-ZERO-BATCH-CAL",
            manufacturer="Maker D",
            unit="次",
            package_unit="盒",
            package_to_min_qty=1,
            warning_days=5,
            warning_stock=1,
            follow_instrument=False,
            is_calibrator=True,
        )
        db.add_all([reagent_expired, reagent_empty, reagent_normal, reagent_zero_batch_calibrator])
        db.flush()

        expired_batch = BatchStock(
            reagent_id=reagent_expired.id,
            batch_no="DASH-EXP-001",
            expiry_date=date.today().fromordinal(date.today().toordinal() - 1),
            quantity=2,
            loose_min_quantity=0,
            package_unit="盒",
            min_unit="瓶",
            package_to_min_qty=1,
        )
        empty_batch = BatchStock(
            reagent_id=reagent_empty.id,
            batch_no="DASH-EMP-001",
            expiry_date=date.today().fromordinal(date.today().toordinal() + 20),
            quantity=0,
            loose_min_quantity=0,
            package_unit="盒",
            min_unit="瓶",
            package_to_min_qty=1,
        )
        normal_batch = BatchStock(
            reagent_id=reagent_normal.id,
            batch_no="DASH-NOR-001",
            expiry_date=date.today().fromordinal(date.today().toordinal() + 40),
            quantity=5,
            loose_min_quantity=0,
            package_unit="盒",
            min_unit="瓶",
            package_to_min_qty=1,
        )
        zero_calibrator_batch = BatchStock(
            reagent_id=reagent_zero_batch_calibrator.id,
            batch_no="DASH-ZERO-CAL-001",
            expiry_date=date.today().fromordinal(date.today().toordinal() + 40),
            quantity=0,
            loose_min_quantity=0,
            package_unit="盒",
            min_unit="次",
            package_to_min_qty=1,
        )
        stocked_calibrator_batch = BatchStock(
            reagent_id=reagent_zero_batch_calibrator.id,
            batch_no="DASH-ZERO-CAL-002",
            expiry_date=date.today().fromordinal(date.today().toordinal() + 40),
            quantity=10,
            loose_min_quantity=0,
            package_unit="盒",
            min_unit="次",
            package_to_min_qty=1,
        )
        db.add_all([expired_batch, empty_batch, normal_batch, zero_calibrator_batch, stocked_calibrator_batch])
        db.flush()

        db.add(Transaction(
            reagent_id=reagent_expired.id,
            batch_stock_id=expired_batch.id,
            type="OUT",
            quantity=1,
            operator="陈茹",
            source="scan",
            raw_barcode="DASH-EXPIRED",
            package_unit="盒",
            min_unit="瓶",
            package_to_min_qty=1,
            remark=None,
        ))
        db.commit()

    response = client.get("/api/dashboard/overview")
    assert response.status_code == 200
    payload = response.json()

    assert payload["pending_tasks"]["expired_batch_count"] == 1
    assert payload["pending_tasks"]["low_stock_batch_count"] == 3
    assert payload["pending_tasks"]["low_or_empty_reagent_count"] == 1
    assert payload["pending_tasks"]["zero_stock_unarchived_count"] == 1

    warning_names = [item["reagent_name"] for item in payload["warning_reagents"]]
    assert warning_names[:2] == ["DASH-EXPIRED", "DASH-EMPTY"]
    assert "DASH-ZERO-BATCH-CAL" in warning_names

    assert len(payload["recent_outbounds"]) == 1
    assert payload["recent_outbounds"][0]["type"] == "OUT"
    assert payload["recent_outbounds"][0]["reagent_name"] == "DASH-EXPIRED"


def test_mobile_overview_returns_compact_recent_outbounds(tmp_path):
    client, SessionLocal = build_client(tmp_path)
    _, instrument_id = get_first_catalog_session(SessionLocal)

    with SessionLocal() as db:
        reagent = Reagent(
            instrument_id=instrument_id,
            name="MOBILE-OVERVIEW",
            match_mode="exact",
            match_pattern="MOBILE-OVERVIEW",
            raw_barcode_sample="MOBILE-OVERVIEW",
            manufacturer="Maker M",
            unit="瓶",
            package_unit="盒",
            package_to_min_qty=1,
            warning_days=30,
            warning_stock=5,
            follow_instrument=False,
        )
        db.add(reagent)
        db.flush()

        batch = BatchStock(
            reagent_id=reagent.id,
            batch_no="MOBILE-BATCH-001",
            expiry_date=date.today().fromordinal(date.today().toordinal() + 7),
            quantity=2,
            loose_min_quantity=0,
            package_unit="盒",
            min_unit="瓶",
            package_to_min_qty=1,
        )
        db.add(batch)
        db.flush()

        for index in range(6):
            db.add(Transaction(
                reagent_id=reagent.id,
                batch_stock_id=batch.id,
                type="OUT",
                quantity=1,
                operator="陈茽",
                source="manual",
                raw_barcode=f"MOBILE-OVERVIEW-{index}",
                package_unit="盒",
                min_unit="瓶",
                package_to_min_qty=1,
                remark=None,
                created_at=datetime(2026, 4, 8, 8, index, 0),
            ))
        db.commit()

    response = client.get("/api/mobile/overview")
    assert response.status_code == 200
    payload = response.json()

    assert payload["pending_tasks"]["expiring_reagent_count"] >= 1
    assert len(payload["recent_outbounds"]) == 5
    assert payload["recent_outbounds"][0]["reagent_name"] == "MOBILE-OVERVIEW"
    assert len(payload["warning_reagents"]) <= 4


def test_mobile_overview_filters_by_instrument(tmp_path):
    client, SessionLocal = build_client(tmp_path)

    with SessionLocal() as db:
        work_group = db.query(WorkGroup).order_by(WorkGroup.id.asc()).first()
        first_instrument = (
            db.query(Instrument)
            .filter(Instrument.work_group_id == work_group.id)
            .order_by(Instrument.id.asc())
            .first()
        )
        second_instrument = Instrument(
            work_group_id=work_group.id,
            name="SECOND-INSTRUMENT",
            model="MODEL-B",
            default_warning_days=30,
            default_warning_stock=5,
        )
        db.add(second_instrument)
        db.flush()

        first_reagent = Reagent(
            instrument_id=first_instrument.id,
            name="FILTERED-A",
            match_mode="exact",
            match_pattern="FILTERED-A",
            manufacturer="Maker A",
            unit="瓶",
            package_unit="盒",
            package_to_min_qty=1,
            warning_days=30,
            warning_stock=5,
            follow_instrument=False,
        )
        second_reagent = Reagent(
            instrument_id=second_instrument.id,
            name="FILTERED-B",
            match_mode="exact",
            match_pattern="FILTERED-B",
            manufacturer="Maker B",
            unit="瓶",
            package_unit="盒",
            package_to_min_qty=1,
            warning_days=30,
            warning_stock=5,
            follow_instrument=False,
        )
        db.add_all([first_reagent, second_reagent])
        db.flush()

        first_batch = BatchStock(
            reagent_id=first_reagent.id,
            batch_no="FILTER-A-001",
            expiry_date=date.today().fromordinal(date.today().toordinal() + 5),
            quantity=2,
            loose_min_quantity=0,
            package_unit="盒",
            min_unit="瓶",
            package_to_min_qty=1,
        )
        second_batch = BatchStock(
            reagent_id=second_reagent.id,
            batch_no="FILTER-B-001",
            expiry_date=date.today().fromordinal(date.today().toordinal() + 5),
            quantity=2,
            loose_min_quantity=0,
            package_unit="盒",
            min_unit="瓶",
            package_to_min_qty=1,
        )
        db.add_all([first_batch, second_batch])
        db.flush()

        db.add_all([
            Transaction(
                reagent_id=first_reagent.id,
                batch_stock_id=first_batch.id,
                type="OUT",
                quantity=1,
                operator="闄堣尳",
                source="manual",
                raw_barcode="FILTER-A",
                package_unit="盒",
                min_unit="瓶",
                package_to_min_qty=1,
                remark=None,
                created_at=datetime(2026, 4, 9, 8, 0, 0),
            ),
            Transaction(
                reagent_id=second_reagent.id,
                batch_stock_id=second_batch.id,
                type="OUT",
                quantity=1,
                operator="闄堣尳",
                source="manual",
                raw_barcode="FILTER-B",
                package_unit="盒",
                min_unit="瓶",
                package_to_min_qty=1,
                remark=None,
                created_at=datetime(2026, 4, 9, 9, 0, 0),
            ),
        ])
        db.commit()
        first_instrument_id = first_instrument.id
        first_instrument_name = first_instrument.name

    filtered_response = client.get(f"/api/mobile/overview?instrument_id={first_instrument_id}")
    assert filtered_response.status_code == 200
    filtered_payload = filtered_response.json()

    assert filtered_payload["pending_tasks"]["expiring_reagent_count"] == 1
    assert [item["reagent_name"] for item in filtered_payload["warning_reagents"]] == ["FILTERED-A"]
    assert [item["instrument_name"] for item in filtered_payload["recent_outbounds"]] == [first_instrument_name]
    assert [item["reagent_name"] for item in filtered_payload["recent_outbounds"]] == ["FILTERED-A"]


def test_mobile_qr_returns_svg(tmp_path):
    client, _ = build_client(tmp_path)

    response = client.get("/api/mobile/qr", params={"target": "http://127.0.0.1:8000/m"})
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("image/svg+xml")
    assert "<svg" in response.text


def test_reagent_status_filter_and_restore(tmp_path):
    client, SessionLocal = build_client(tmp_path)
    _, instrument_id = get_first_catalog_session(SessionLocal)

    create_response = client.post("/api/reagents", json={
        "instrument_id": instrument_id,
        "name": "可恢复归档试剂",
        "match_mode": "exact",
        "match_pattern": "RESTORE-REAGENT-001",
        "manufacturer": "测试厂家",
        "unit": "瓶",
        "package_unit": "盒",
        "package_to_min_qty": 1,
        "warning_days": 20,
        "warning_stock": 3,
    })
    reagent_id = create_response.json()["id"]

    archive_response = client.post(f"/api/reagents/{reagent_id}/archive")
    assert archive_response.status_code == 200

    active_response = client.get("/api/reagents?status=active")
    assert all(item["id"] != reagent_id for item in active_response.json())

    archived_response = client.get("/api/reagents?status=archived")
    assert any(item["id"] == reagent_id for item in archived_response.json())

    restore_response = client.post(f"/api/reagents/{reagent_id}/restore")
    assert restore_response.status_code == 200

    active_response = client.get("/api/reagents?status=active")
    assert any(item["id"] == reagent_id for item in active_response.json())


def test_unit_option_crud_and_scope_filter(tmp_path):
    client, _ = build_client(tmp_path)

    list_response = client.get("/api/unit-options")
    assert list_response.status_code == 200
    payload = list_response.json()
    assert any(item["name"] == "盒" for item in payload)
    assert all(item["is_active"] is True for item in payload)

    create_response = client.post("/api/unit-options", json={
        "name": "板",
        "scope": "package",
        "is_active": True,
        "sort_order": 5,
    })
    assert create_response.status_code == 201
    option_id = create_response.json()["id"]

    package_response = client.get("/api/unit-options?scope=package")
    assert any(item["name"] == "板" for item in package_response.json())

    update_response = client.put(f"/api/unit-options/{option_id}", json={
        "name": "板",
        "scope": "both",
        "is_active": False,
        "sort_order": 6,
    })
    assert update_response.status_code == 200
    assert update_response.json()["is_active"] is False

    active_response = client.get("/api/unit-options")
    assert all(item["id"] != option_id for item in active_response.json())

    all_response = client.get("/api/unit-options?include_inactive=true")
    assert any(item["id"] == option_id and item["is_active"] is False for item in all_response.json())


def test_clear_stock_keeps_reagents_and_parse_logs(tmp_path):
    client, SessionLocal = build_client(tmp_path)
    _, instrument_id = get_first_catalog_session(SessionLocal)

    with SessionLocal() as db:
        reagent = Reagent(
            instrument_id=instrument_id,
            name="清库存保留档案试剂",
            match_mode="exact",
            match_pattern="CLEAR-STOCK-001",
            raw_barcode_sample="CLEAR-STOCK-001",
            manufacturer="测试厂家",
            unit="瓶",
            package_unit="盒",
            package_to_min_qty=1,
            warning_days=30,
            warning_stock=5,
        )
        db.add(reagent)
        db.flush()

        batch = BatchStock(
            reagent_id=reagent.id,
            batch_no="CLEAR-STOCK-BATCH",
            expiry_date=date(2031, 1, 1),
            quantity=3,
            package_unit="盒",
            min_unit="瓶",
            package_to_min_qty=1,
        )
        db.add(batch)
        db.flush()

        db.add(ParseLog(raw_barcode="CLEAR-STOCK-001", match_result="exact", matched_reagent_id=reagent.id))
        db.add(Transaction(
            reagent_id=reagent.id,
            batch_stock_id=batch.id,
            type="IN",
            quantity=3,
            operator="陈芽",
            source="manual",
            raw_barcode="CLEAR-STOCK-001",
            package_unit="盒",
            min_unit="瓶",
            package_to_min_qty=1,
            remark=None,
        ))
        db.commit()
        reagent_id = reagent.id

    response = client.post("/api/admin/clear-stock", json={"password": "wtsyyjyk"})
    assert response.status_code == 200
    assert response.json()["cleared_batch_count"] == 1

    with SessionLocal() as db:
        assert db.query(Reagent).filter(Reagent.id == reagent_id).count() == 1
        assert db.query(ParseLog).filter(ParseLog.matched_reagent_id == reagent_id).count() == 1
        assert db.query(BatchStock).count() == 0
        tx = db.query(Transaction).filter(Transaction.reagent_id == reagent_id).first()
        assert tx is not None
        assert tx.batch_stock_id is None


def test_clear_transactions_with_filters_only_removes_matching_rows(tmp_path):
    client, SessionLocal = build_client(tmp_path)
    work_group_id, instrument_id = get_first_catalog_session(SessionLocal)

    with SessionLocal() as db:
        reagent = Reagent(
            instrument_id=instrument_id,
            name="筛选清流水试剂",
            match_mode="exact",
            match_pattern="CLEAR-TX-001",
            raw_barcode_sample="CLEAR-TX-001",
            manufacturer="测试厂家",
            unit="瓶",
            package_unit="盒",
            package_to_min_qty=1,
            warning_days=30,
            warning_stock=5,
        )
        db.add(reagent)
        db.flush()

        batch = BatchStock(
            reagent_id=reagent.id,
            batch_no="CLEAR-TX-BATCH",
            expiry_date=date(2031, 1, 1),
            quantity=1,
            package_unit="盒",
            min_unit="瓶",
            package_to_min_qty=1,
        )
        db.add(batch)
        db.flush()

        keep_reagent = Reagent(
            instrument_id=instrument_id,
            name="保留流水试剂",
            match_mode="exact",
            match_pattern="KEEP-TX-001",
            raw_barcode_sample="KEEP-TX-001",
            manufacturer="测试厂家",
            unit="瓶",
            package_unit="盒",
            package_to_min_qty=1,
            warning_days=30,
            warning_stock=5,
        )
        db.add(keep_reagent)
        db.flush()

        db.add_all([
            Transaction(
                reagent_id=reagent.id,
                batch_stock_id=batch.id,
                type="IN",
                quantity=1,
                operator="陈芽",
                source="manual",
                raw_barcode="CLEAR-TX-001",
                package_unit="盒",
                min_unit="瓶",
                package_to_min_qty=1,
                remark=None,
                created_at=datetime(2026, 4, 1, 9, 0, 0),
            ),
            Transaction(
                reagent_id=keep_reagent.id,
                batch_stock_id=None,
                type="OUT",
                quantity=1,
                operator="陈芽",
                source="manual",
                raw_barcode="KEEP-TX-001",
                package_unit="盒",
                min_unit="瓶",
                package_to_min_qty=1,
                remark=None,
                created_at=datetime(2026, 4, 5, 9, 0, 0),
            ),
        ])
        calibration_transaction = Transaction(
            reagent_id=reagent.id,
            batch_stock_id=batch.id,
            type="CALIBRATION",
            quantity=1,
            operator="operator",
            source="scan",
            raw_barcode="CLEAR-CAL-001",
            package_unit="box",
            min_unit="bottle",
            package_to_min_qty=1,
            remark=None,
            created_at=datetime(2026, 4, 1, 10, 0, 0),
        )
        db.add(calibration_transaction)
        db.flush()
        db.add_all([
            CalibrationRecord(
                calibrator_reagent_id=reagent.id,
                calibrator_batch_stock_id=batch.id,
                target_reagent_id=reagent.id,
                transaction_id=calibration_transaction.id,
                target_old_batch_no="OLD-CLEAR",
                target_new_batch_no="NEW-CLEAR",
                usage_times=1,
                result_unit="nmmol/L",
                operator="operator",
                created_at=datetime(2026, 4, 1, 10, 0, 0),
            ),
            CalibrationRecord(
                calibrator_reagent_id=reagent.id,
                calibrator_batch_stock_id=batch.id,
                target_reagent_id=reagent.id,
                transaction_id=None,
                target_old_batch_no="OLD-ORPHAN",
                target_new_batch_no="NEW-ORPHAN",
                usage_times=1,
                result_unit="nmmol/L",
                operator="operator",
                created_at=datetime(2026, 4, 1, 11, 0, 0),
            ),
            CalibrationRecord(
                calibrator_reagent_id=keep_reagent.id,
                calibrator_batch_stock_id=batch.id,
                target_reagent_id=keep_reagent.id,
                transaction_id=None,
                target_old_batch_no="OLD-KEEP",
                target_new_batch_no="NEW-KEEP",
                usage_times=1,
                result_unit="nmmol/L",
                operator="operator",
                created_at=datetime(2026, 4, 5, 10, 0, 0),
            ),
        ])
        db.commit()
        reagent_id = reagent.id
        keep_reagent_id = keep_reagent.id

    response = client.post("/api/admin/clear-transactions", json={
        "password": "wtsyyjyk",
        "work_group_id": work_group_id,
        "instrument_id": instrument_id,
        "reagent_id": reagent_id,
        "start_date": "2026-04-01",
        "end_date": "2026-04-02",
    })
    assert response.status_code == 200
    assert response.json()["deleted_count"] == 2
    assert response.json()["deleted_calibration_record_count"] == 2

    with SessionLocal() as db:
        assert db.query(Transaction).filter(Transaction.reagent_id == reagent_id).count() == 0
        assert db.query(Transaction).filter(Transaction.reagent_id == keep_reagent_id).count() == 1
        assert db.query(CalibrationRecord).filter(CalibrationRecord.target_reagent_id == reagent_id).count() == 0
        assert db.query(CalibrationRecord).filter(CalibrationRecord.target_reagent_id == keep_reagent_id).count() == 1
        assert db.query(BatchStock).count() == 1


def test_stock_adjustment_access_and_apply_updates_batch_totals(tmp_path):
    client, SessionLocal = build_client(tmp_path)
    _, instrument_id = get_first_catalog_session(SessionLocal)

    create_group_response = client.post("/api/reagent-groups", json={
        "name": "快速调整组",
        "sort_order": 5,
        "is_active": True,
    })
    assert create_group_response.status_code == 201
    reagent_group_id = create_group_response.json()["id"]

    create_reagent_response = client.post("/api/reagents", json={
        "instrument_id": instrument_id,
        "name": "库存调整试剂",
        "match_mode": "exact",
        "match_pattern": "STOCK-ADJUST-001",
        "manufacturer": "维护厂家",
        "unit": "瓶",
        "package_unit": "盒",
        "package_to_min_qty": 4,
        "warning_days": 30,
        "warning_stock": 5,
        "reagent_group_id": reagent_group_id,
    })
    assert create_reagent_response.status_code == 201
    reagent_id = create_reagent_response.json()["id"]

    with SessionLocal() as db:
        db.add_all([
            BatchStock(
                reagent_id=reagent_id,
                batch_no="ADJUST-BATCH-1",
                expiry_date=date(2031, 1, 1),
                quantity=2,
                loose_min_quantity=1,
                package_unit="盒",
                min_unit="瓶",
                package_to_min_qty=4,
            ),
            BatchStock(
                reagent_id=reagent_id,
                batch_no="ADJUST-BATCH-2",
                expiry_date=date(2031, 6, 1),
                quantity=1,
                loose_min_quantity=0,
                package_unit="盒",
                min_unit="瓶",
                package_to_min_qty=4,
            ),
        ])
        db.commit()
        batches = (
            db.query(BatchStock)
            .filter(BatchStock.reagent_id == reagent_id)
            .order_by(BatchStock.batch_no.asc())
            .all()
        )
        first_batch_id = batches[0].id
        second_batch_id = batches[1].id

    access_response = client.post("/api/admin/stock-adjustments/access", json={"password": "wtsyyjyk"})
    assert access_response.status_code == 200
    payload = access_response.json()
    reagent_payload = next(
        reagent
        for group in payload
        for instrument in group["instruments"]
        for reagent in instrument["reagents"]
        if reagent["id"] == reagent_id
    )
    assert reagent_payload["reagent_group_id"] == reagent_group_id
    assert [batch["batch_no"] for batch in reagent_payload["batches"]] == ["ADJUST-BATCH-1", "ADJUST-BATCH-2"]

    apply_response = client.post("/api/admin/stock-adjustments/apply", json={
        "password": "wtsyyjyk",
        "items": [
            {"batch_stock_id": first_batch_id, "target_min_quantity": 6},
            {"batch_stock_id": second_batch_id, "target_min_quantity": 3},
        ],
    })
    assert apply_response.status_code == 200
    assert apply_response.json()["updated_batch_count"] == 2
    assert apply_response.json()["updated_reagent_count"] == 1

    with SessionLocal() as db:
        first_batch = db.query(BatchStock).filter(BatchStock.id == first_batch_id).first()
        second_batch = db.query(BatchStock).filter(BatchStock.id == second_batch_id).first()
        assert first_batch.quantity == 1
        assert first_batch.loose_min_quantity == 2
        assert second_batch.quantity == 0
        assert second_batch.loose_min_quantity == 3


def test_dangerous_admin_actions_require_password(tmp_path):
    client, _ = build_client(tmp_path)

    clear_stock = client.post("/api/admin/clear-stock", json={"password": "wrong"})
    assert clear_stock.status_code == 403

    clear_transactions = client.post("/api/admin/clear-transactions", json={"password": "wrong"})
    assert clear_transactions.status_code == 403

    stock_adjustment_access = client.post("/api/admin/stock-adjustments/access", json={"password": "wrong"})
    assert stock_adjustment_access.status_code == 403

    stock_adjustment_apply = client.post("/api/admin/stock-adjustments/apply", json={"password": "wrong", "items": []})
    assert stock_adjustment_apply.status_code == 403

    reset_response = client.post("/api/admin/reset-reagent-data", json={"password": "wrong"})
    assert reset_response.status_code == 403


def test_inventory_keeps_zero_stock_reagent_and_marks_empty(tmp_path):
    client, SessionLocal = build_client(tmp_path)
    _, instrument_id = get_first_catalog_session(SessionLocal)

    with SessionLocal() as db:
        reagent = Reagent(
            instrument_id=instrument_id,
            name="零库存试剂",
            match_mode="exact",
            match_pattern="EMPTY-001",
            raw_barcode_sample="EMPTY-001",
            manufacturer="测试厂家",
            unit="瓶",
            package_unit="盒",
            package_to_min_qty=1,
            warning_days=10,
            warning_stock=2,
            follow_instrument=False,
        )
        db.add(reagent)
        db.flush()

        db.add(BatchStock(
            reagent_id=reagent.id,
            batch_no="EMPTY-BATCH-1",
            expiry_date=date.today().fromordinal(date.today().toordinal() + 30),
            quantity=0,
            loose_min_quantity=0,
            package_unit="盒",
            min_unit="瓶",
            package_to_min_qty=1,
        ))
        db.commit()
        reagent_id = reagent.id

    inventory_response = client.get("/api/inventory")
    assert inventory_response.status_code == 200
    reagent_payload = next(
        reagent
        for group in inventory_response.json()
        for instrument in group["instruments"]
        for reagent in instrument["reagents"]
        if reagent["id"] == reagent_id
    )
    assert reagent_payload["total_min_quantity"] == 0
    assert reagent_payload["status_level"] == "empty"
    assert reagent_payload["batch_count"] == 1


def test_archive_zero_stock_reagent_hides_it_from_inventory(tmp_path):
    client, SessionLocal = build_client(tmp_path)
    _, instrument_id = get_first_catalog_session(SessionLocal)

    with SessionLocal() as db:
        reagent = Reagent(
            instrument_id=instrument_id,
            name="待归档零库存试剂",
            match_mode="exact",
            match_pattern="ARCHIVE-EMPTY-001",
            raw_barcode_sample="ARCHIVE-EMPTY-001",
            manufacturer="测试厂家",
            unit="瓶",
            package_unit="盒",
            package_to_min_qty=1,
            warning_days=10,
            warning_stock=2,
            follow_instrument=False,
        )
        db.add(reagent)
        db.flush()

        db.add(BatchStock(
            reagent_id=reagent.id,
            batch_no="ARCHIVE-EMPTY-BATCH",
            expiry_date=date.today().fromordinal(date.today().toordinal() + 30),
            quantity=0,
            loose_min_quantity=0,
            package_unit="盒",
            min_unit="瓶",
            package_to_min_qty=1,
        ))
        db.commit()
        reagent_id = reagent.id

    archive_response = client.post(f"/api/reagents/{reagent_id}/archive")
    assert archive_response.status_code == 200
    assert archive_response.json()["success"] is True

    with SessionLocal() as db:
        reagent = db.query(Reagent).filter(Reagent.id == reagent_id).first()
        assert reagent.is_archived is True

    inventory_response = client.get("/api/inventory")
    assert inventory_response.status_code == 200
    reagent_ids = [
        reagent["id"]
        for group in inventory_response.json()
        for instrument in group["instruments"]
        for reagent in instrument["reagents"]
    ]
    assert reagent_id not in reagent_ids


def test_bulk_archive_zero_stock_reagents_only_archives_empty_ones(tmp_path):
    client, SessionLocal = build_client(tmp_path)
    _, instrument_id = get_first_catalog_session(SessionLocal)

    with SessionLocal() as db:
        empty_reagent = Reagent(
            instrument_id=instrument_id,
            name="批量归档零库存试剂",
            match_mode="exact",
            match_pattern="ARCHIVE-BULK-EMPTY-001",
            raw_barcode_sample="ARCHIVE-BULK-EMPTY-001",
            manufacturer="测试厂家",
            unit="瓶",
            package_unit="盒",
            package_to_min_qty=1,
            warning_days=10,
            warning_stock=2,
            follow_instrument=False,
        )
        stocked_reagent = Reagent(
            instrument_id=instrument_id,
            name="批量归档保留试剂",
            match_mode="exact",
            match_pattern="ARCHIVE-BULK-STOCK-001",
            raw_barcode_sample="ARCHIVE-BULK-STOCK-001",
            manufacturer="测试厂家",
            unit="瓶",
            package_unit="盒",
            package_to_min_qty=1,
            warning_days=10,
            warning_stock=2,
            follow_instrument=False,
        )
        db.add_all([empty_reagent, stocked_reagent])
        db.flush()

        db.add(BatchStock(
            reagent_id=empty_reagent.id,
            batch_no="ARCHIVE-BULK-EMPTY-BATCH",
            expiry_date=date.today().fromordinal(date.today().toordinal() + 30),
            quantity=0,
            loose_min_quantity=0,
            package_unit="盒",
            min_unit="瓶",
            package_to_min_qty=1,
        ))
        db.add(BatchStock(
            reagent_id=stocked_reagent.id,
            batch_no="ARCHIVE-BULK-STOCK-BATCH",
            expiry_date=date.today().fromordinal(date.today().toordinal() + 30),
            quantity=3,
            loose_min_quantity=0,
            package_unit="盒",
            min_unit="瓶",
            package_to_min_qty=1,
        ))
        db.commit()
        empty_reagent_id = empty_reagent.id
        stocked_reagent_id = stocked_reagent.id

    archive_response = client.post("/api/reagents/archive-zero-stock")
    assert archive_response.status_code == 200
    payload = archive_response.json()
    assert payload["success"] is True
    assert empty_reagent_id in payload["reagent_ids"]
    assert stocked_reagent_id not in payload["reagent_ids"]

    with SessionLocal() as db:
        empty_reagent = db.query(Reagent).filter(Reagent.id == empty_reagent_id).first()
        stocked_reagent = db.query(Reagent).filter(Reagent.id == stocked_reagent_id).first()
        assert empty_reagent.is_archived is True
        assert stocked_reagent.is_archived is False


def test_archive_zero_stock_reagents_archives_empty_batches_only(tmp_path):
    client, SessionLocal = build_client(tmp_path)
    _, instrument_id = get_first_catalog_session(SessionLocal)

    with SessionLocal() as db:
        reagent = Reagent(
            instrument_id=instrument_id,
            name="TT3",
            match_mode="exact",
            match_pattern="ARCHIVE-BATCH-TT3-001",
            raw_barcode_sample="ARCHIVE-BATCH-TT3-001",
            manufacturer="TEST",
            unit="瓶",
            package_unit="盒",
            package_to_min_qty=1,
            warning_days=30,
            warning_stock=5,
            follow_instrument=False,
        )
        db.add(reagent)
        db.flush()

        empty_batch = BatchStock(
            reagent_id=reagent.id,
            batch_no="538621",
            expiry_date=date(2031, 7, 31),
            quantity=0,
            loose_min_quantity=0,
            package_unit="盒",
            min_unit="瓶",
            package_to_min_qty=1,
        )
        stocked_batch = BatchStock(
            reagent_id=reagent.id,
            batch_no="538990",
            expiry_date=date(2031, 9, 30),
            quantity=3,
            loose_min_quantity=0,
            package_unit="盒",
            min_unit="瓶",
            package_to_min_qty=1,
        )
        db.add_all([empty_batch, stocked_batch])
        db.commit()
        reagent_id = reagent.id
        empty_batch_id = empty_batch.id

    archive_response = client.post("/api/reagents/archive-zero-stock")
    assert archive_response.status_code == 200
    payload = archive_response.json()
    assert payload["success"] is True
    assert reagent_id not in payload["reagent_ids"]
    assert empty_batch_id in payload["batch_stock_ids"]
    assert payload["archived_batch_count"] >= 1

    inventory_response = client.get("/api/inventory")
    assert inventory_response.status_code == 200
    reagent_payload = next(
        reagent
        for group in inventory_response.json()
        for instrument in group["instruments"]
        for reagent in instrument["reagents"]
        if reagent["id"] == reagent_id
    )
    assert [batch["batch_no"] for batch in reagent_payload["batches"]] == ["538990"]

    with SessionLocal() as db:
        empty_batch = db.query(BatchStock).filter(BatchStock.id == empty_batch_id).first()
        reagent = db.query(Reagent).filter(Reagent.id == reagent_id).first()
        assert reagent.is_archived is False
        assert empty_batch.is_archived is True


def test_archived_reagent_can_be_scanned_and_reactivated_by_inbound(tmp_path):
    client, SessionLocal = build_client(tmp_path)
    _, instrument_id = get_first_catalog_session(SessionLocal)

    with SessionLocal() as db:
        operator_name = db.query(Operator).filter(Operator.is_active.is_(True)).order_by(Operator.id.asc()).first().name
        reagent = Reagent(
            instrument_id=instrument_id,
            name="CA125",
            match_mode="exact",
            match_pattern="SCAN-ARCHIVE-001",
            raw_barcode_sample="SCAN-ARCHIVE-001",
            manufacturer="测试厂家",
            unit="瓶",
            package_unit="盒",
            package_to_min_qty=2,
            warning_days=30,
            warning_stock=5,
            follow_instrument=False,
            is_archived=True,
        )
        db.add(reagent)
        db.commit()
        reagent_id = reagent.id

    parse_response = client.post("/api/barcode/parse", json={"barcode": "SCAN-ARCHIVE-001"})
    assert parse_response.status_code == 200
    parse_payload = parse_response.json()
    assert parse_payload["reagent_id"] == reagent_id
    assert parse_payload["reagent_name"] == "CA125"
    assert parse_payload["match_type"] == "exact"

    inbound_response = client.post("/api/inbound/batch", json={
        "operator": operator_name,
        "items": [{
            "source": "scan",
            "reagent_id": reagent_id,
            "reagent_name": "CA125",
            "batch_no": "SCAN-ARCHIVE-BATCH",
            "expiry_date": "2031-12-31",
            "quantity": 3,
            "package_unit": "盒",
            "min_unit": "瓶",
            "package_to_min_qty": 2,
            "raw_barcode": "SCAN-ARCHIVE-001"
        }]
    })
    assert inbound_response.status_code == 200

    with SessionLocal() as db:
        reagent = db.query(Reagent).filter(Reagent.id == reagent_id).first()
        assert reagent.is_archived is False
        batch = db.query(BatchStock).filter(BatchStock.reagent_id == reagent_id).first()
        assert batch is not None
        assert batch.quantity == 3


def test_export_inventory_all_work_groups_creates_group_sheets(tmp_path):
    client, SessionLocal = build_client(tmp_path)

    with SessionLocal() as db:
        work_groups = db.query(WorkGroup).order_by(WorkGroup.id.asc()).all()
        first_group = work_groups[0]
        second_group = work_groups[1]
        first_instrument = (
            db.query(Instrument)
            .filter(Instrument.work_group_id == first_group.id)
            .order_by(Instrument.id.asc())
            .first()
        )
        second_instrument = (
            db.query(Instrument)
            .filter(Instrument.work_group_id == second_group.id)
            .order_by(Instrument.id.asc())
            .first()
        )

        reagent_one = Reagent(
            instrument_id=first_instrument.id,
            name="EXPORT-ALL-A",
            match_mode="exact",
            match_pattern="EXPORT-ALL-A",
            raw_barcode_sample="EXPORT-ALL-A",
            manufacturer="Maker A",
            unit="瓶",
            package_unit="盒",
            package_to_min_qty=2,
            warning_days=30,
            warning_stock=5,
            follow_instrument=False,
        )
        reagent_two = Reagent(
            instrument_id=second_instrument.id,
            name="EXPORT-ALL-B",
            match_mode="exact",
            match_pattern="EXPORT-ALL-B",
            raw_barcode_sample="EXPORT-ALL-B",
            manufacturer="Maker B",
            unit="瓶",
            package_unit="盒",
            package_to_min_qty=3,
            warning_days=30,
            warning_stock=5,
            follow_instrument=False,
        )
        archived_reagent = Reagent(
            instrument_id=first_instrument.id,
            name="EXPORT-ARCHIVED",
            match_mode="exact",
            match_pattern="EXPORT-ARCHIVED",
            raw_barcode_sample="EXPORT-ARCHIVED",
            manufacturer="Maker X",
            unit="瓶",
            package_unit="盒",
            package_to_min_qty=1,
            warning_days=30,
            warning_stock=5,
            follow_instrument=False,
            is_archived=True,
        )
        db.add_all([reagent_one, reagent_two, archived_reagent])
        db.flush()

        db.add(BatchStock(
            reagent_id=reagent_one.id,
            batch_no="BATCH-A1",
            expiry_date=date.today().fromordinal(date.today().toordinal() + 15),
            quantity=4,
            loose_min_quantity=1,
            package_unit="盒",
            min_unit="瓶",
            package_to_min_qty=2,
        ))
        db.add(BatchStock(
            reagent_id=reagent_two.id,
            batch_no="BATCH-B1",
            expiry_date=date.today().fromordinal(date.today().toordinal() + 25),
            quantity=2,
            loose_min_quantity=0,
            package_unit="盒",
            min_unit="瓶",
            package_to_min_qty=3,
        ))
        db.add(BatchStock(
            reagent_id=archived_reagent.id,
            batch_no="BATCH-X1",
            expiry_date=date.today().fromordinal(date.today().toordinal() + 25),
            quantity=1,
            loose_min_quantity=0,
            package_unit="盒",
            min_unit="瓶",
            package_to_min_qty=1,
        ))
        db.commit()
        group_names = [group.name for group in work_groups]

    response = client.get("/api/inventory/export")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment;" in response.headers["content-disposition"]

    workbook = load_workbook(filename=BytesIO(response.content))
    assert workbook.sheetnames == group_names

    first_sheet_values = list(workbook[group_names[0]].iter_rows(values_only=True))
    second_sheet_values = list(workbook[group_names[1]].iter_rows(values_only=True))
    assert first_sheet_values[0][0] == "工作组"
    assert any(row[3] == "EXPORT-ALL-A" for row in first_sheet_values[1:])
    assert any(row[3] == "EXPORT-ALL-B" for row in second_sheet_values[1:])
    assert all(row[3] != "EXPORT-ARCHIVED" for row in first_sheet_values[1:])


def test_export_inventory_by_work_group_only_contains_that_group(tmp_path):
    client, SessionLocal = build_client(tmp_path)

    with SessionLocal() as db:
        work_groups = db.query(WorkGroup).order_by(WorkGroup.id.asc()).all()
        target_group = work_groups[0]
        other_group = work_groups[1]
        target_group_id = target_group.id
        target_group_name = target_group.name
        target_instruments = (
            db.query(Instrument)
            .filter(Instrument.work_group_id == target_group.id)
            .order_by(Instrument.id.asc())
            .limit(2)
            .all()
        )
        other_instrument = (
            db.query(Instrument)
            .filter(Instrument.work_group_id == other_group.id)
            .order_by(Instrument.id.asc())
            .first()
        )

        reagent_one = Reagent(
            instrument_id=target_instruments[0].id,
            name="EXPORT-GROUP-A",
            match_mode="exact",
            match_pattern="EXPORT-GROUP-A",
            raw_barcode_sample="EXPORT-GROUP-A",
            manufacturer="Maker A",
            unit="瓶",
            package_unit="盒",
            package_to_min_qty=1,
            warning_days=30,
            warning_stock=5,
            follow_instrument=False,
        )
        reagent_two = Reagent(
            instrument_id=target_instruments[1].id,
            name="EXPORT-GROUP-B",
            match_mode="exact",
            match_pattern="EXPORT-GROUP-B",
            raw_barcode_sample="EXPORT-GROUP-B",
            manufacturer="Maker B",
            unit="瓶",
            package_unit="盒",
            package_to_min_qty=1,
            warning_days=30,
            warning_stock=5,
            follow_instrument=False,
        )
        reagent_three = Reagent(
            instrument_id=other_instrument.id,
            name="EXPORT-GROUP-X",
            match_mode="exact",
            match_pattern="EXPORT-GROUP-X",
            raw_barcode_sample="EXPORT-GROUP-X",
            manufacturer="Maker X",
            unit="瓶",
            package_unit="盒",
            package_to_min_qty=1,
            warning_days=30,
            warning_stock=5,
            follow_instrument=False,
        )
        db.add_all([reagent_one, reagent_two, reagent_three])
        db.flush()
        for reagent_id in (reagent_one.id, reagent_two.id, reagent_three.id):
            db.add(BatchStock(
                reagent_id=reagent_id,
                batch_no=f"BATCH-{reagent_id}",
                expiry_date=date.today().fromordinal(date.today().toordinal() + 10),
                quantity=1,
                loose_min_quantity=0,
                package_unit="盒",
                min_unit="瓶",
                package_to_min_qty=1,
            ))
        db.commit()

    response = client.get(f"/api/inventory/export?work_group_id={target_group_id}")
    assert response.status_code == 200

    workbook = load_workbook(filename=BytesIO(response.content))
    assert workbook.sheetnames == [target_group_name]
    sheet_rows = list(workbook[target_group_name].iter_rows(values_only=True))[1:]
    exported_reagents = {row[3] for row in sheet_rows}
    assert exported_reagents == {"EXPORT-GROUP-A", "EXPORT-GROUP-B"}


def test_export_inventory_by_instrument_only_contains_that_instrument(tmp_path):
    client, SessionLocal = build_client(tmp_path)

    with SessionLocal() as db:
        work_group = db.query(WorkGroup).order_by(WorkGroup.id.asc()).first()
        work_group_id = work_group.id
        work_group_name = work_group.name
        instruments = (
            db.query(Instrument)
            .filter(Instrument.work_group_id == work_group.id)
            .order_by(Instrument.id.asc())
            .limit(2)
            .all()
        )
        target_instrument = instruments[0]
        other_instrument = instruments[1]
        target_instrument_id = target_instrument.id
        target_instrument_name = target_instrument.name

        reagent_one = Reagent(
            instrument_id=target_instrument.id,
            name="EXPORT-INSTRUMENT-A",
            match_mode="exact",
            match_pattern="EXPORT-INSTRUMENT-A",
            raw_barcode_sample="EXPORT-INSTRUMENT-A",
            manufacturer="Maker A",
            unit="瓶",
            package_unit="盒",
            package_to_min_qty=1,
            warning_days=30,
            warning_stock=5,
            follow_instrument=False,
        )
        reagent_two = Reagent(
            instrument_id=other_instrument.id,
            name="EXPORT-INSTRUMENT-B",
            match_mode="exact",
            match_pattern="EXPORT-INSTRUMENT-B",
            raw_barcode_sample="EXPORT-INSTRUMENT-B",
            manufacturer="Maker B",
            unit="瓶",
            package_unit="盒",
            package_to_min_qty=1,
            warning_days=30,
            warning_stock=5,
            follow_instrument=False,
        )
        db.add_all([reagent_one, reagent_two])
        db.flush()
        db.add(BatchStock(
            reagent_id=reagent_one.id,
            batch_no="BATCH-IA",
            expiry_date=date.today().fromordinal(date.today().toordinal() + 10),
            quantity=2,
            loose_min_quantity=0,
            package_unit="盒",
            min_unit="瓶",
            package_to_min_qty=1,
        ))
        db.add(BatchStock(
            reagent_id=reagent_two.id,
            batch_no="BATCH-IB",
            expiry_date=date.today().fromordinal(date.today().toordinal() + 20),
            quantity=2,
            loose_min_quantity=0,
            package_unit="盒",
            min_unit="瓶",
            package_to_min_qty=1,
        ))
        db.commit()

    response = client.get(
        f"/api/inventory/export?work_group_id={work_group_id}&instrument_id={target_instrument_id}"
    )
    assert response.status_code == 200

    workbook = load_workbook(filename=BytesIO(response.content))
    assert workbook.sheetnames == [work_group_name]
    sheet_rows = list(workbook[work_group_name].iter_rows(values_only=True))[1:]
    exported_instruments = {row[1] for row in sheet_rows}
    exported_reagents = {row[3] for row in sheet_rows}
    assert exported_instruments == {target_instrument_name}
    assert exported_reagents == {"EXPORT-INSTRUMENT-A"}


def test_manual_inbound_new_reagent_inherits_instrument_defaults(tmp_path):
    client, SessionLocal = build_client(tmp_path)
    work_group_id, instrument_id = get_first_catalog_session(SessionLocal)

    with SessionLocal() as db:
        operator_name = db.query(Operator).filter(Operator.is_active.is_(True)).order_by(Operator.id.asc()).first().name

    with SessionLocal() as db:
        instrument = db.query(Instrument).filter(Instrument.id == instrument_id).first()
        instrument.default_warning_days = 14
        instrument.default_warning_stock = 7
        db.commit()

    response = client.post("/api/inbound/batch", json={
        "operator": operator_name,
        "items": [{
            "source": "manual",
            "work_group_id": work_group_id,
            "instrument_id": instrument_id,
            "reagent_name": "继承仪器默认试剂",
            "manufacturer": "默认值厂家",
            "batch_no": "DEFAULT-001",
            "expiry_date": "2031-12-31",
            "quantity": 1,
            "package_unit": "盒",
            "min_unit": "瓶",
            "package_to_min_qty": 1,
            "raw_barcode": "MANUAL::DEFAULT::001"
        }]
    })
    assert response.status_code == 200

    with SessionLocal() as db:
        reagent = db.query(Reagent).filter(Reagent.name == "继承仪器默认试剂").first()
        assert reagent is not None
        assert reagent.warning_days == 14
        assert reagent.warning_stock == 7
        assert reagent.follow_instrument is False
